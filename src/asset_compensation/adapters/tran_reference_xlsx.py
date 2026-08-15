"""Read-only indexes for the external FA&GL and optional CCDC workbooks."""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from asset_compensation.domain import DepreciationGroup, ReferenceStatus, ValidationError

_FA_SHEETS = {
    "VNG-Asset": (3, 4, "Asset", "VNG"),
    "VNG-Tool": (3, 4, "Tool", "VNG"),
    "VNGS-Asset": (2, 3, "Asset", "VNGS"),
    "VNGS-Tool": (2, 3, "Tool", "VNGS"),
}
_FA_HEADERS = {
    2: "company name",
    7: "cost center",
    8: "product code",
    10: "location",
    12: "asset",
    16: "tag number",
    22: "date of depreciation",
    25: "life(month)",
    27: "cost",
}
_NONPHYSICAL_GROUP_TYPES = {"service", "software", "virtual asset"}


class TranReferenceError(ValidationError):
    """Raised when a reference workbook violates the documented contract."""


def _plain_text(value: object) -> str:
    return str(value or "").strip()


def _normalized_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", _plain_text(value))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def _tag(value: object) -> str:
    return re.sub(r"\s+", "", _plain_text(value)).upper()


def _date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _plain_text(value)
    if not text:
        return None
    for parser in (
        date.fromisoformat,
        lambda raw: datetime.strptime(raw, "%d/%m/%Y").date(),
        lambda raw: datetime.strptime(raw, "%m/%d/%Y").date(),
    ):
        try:
            return parser(text)
        except ValueError:
            continue
    return None


def _whole_vnd(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    text = str(value).strip().replace(",", "")
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if not amount.is_finite() or amount < 0 or amount != amount.to_integral_value():
        return None
    return amount


@dataclass(frozen=True, slots=True)
class FaGlRecord:
    """One exact Tag Number match from a documented FA&GL sheet."""

    tag_number: str
    asset_number: str | None
    start_date: date | None
    cost: Decimal | None
    book: str
    entity: str
    cost_center: str | None
    product_code: str | None
    location: str | None
    company_name: str | None
    source_file: str
    source_sheet: str
    source_row: int

    @property
    def source_note(self) -> str:
        return (
            f"Tra từ file {self.source_file}, sheet {self.source_sheet}, theo Tagnumber "
            "cột P; Asset Number=cột L, Cost=cột AA, Ngày bắt đầu khấu hao=cột V, "
            "Cost center=cột G, Product code=cột H, Location=cột J."
        )


@dataclass(frozen=True, slots=True)
class FaGlLookup:
    """Result of an exact lookup across all four FA&GL sheets."""

    status: ReferenceStatus
    matches: tuple[FaGlRecord, ...]

    @property
    def record(self) -> FaGlRecord | None:
        return self.matches[0] if self.status is ReferenceStatus.MATCHED else None


class FaGlWorkbookIndex:
    """In-memory exact-match index built without modifying the source workbook."""

    def __init__(self, records: dict[str, tuple[FaGlRecord, ...]], source_path: Path) -> None:
        self._records = records
        self.source_path = source_path

    @classmethod
    def from_path(cls, path: str | Path) -> FaGlWorkbookIndex:
        source_path = Path(path).resolve()
        if not source_path.is_file() or source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise TranReferenceError("FA&GL source must be an existing .xlsx or .xlsm file")
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        indexed: dict[str, list[FaGlRecord]] = defaultdict(list)
        try:
            missing = [sheet for sheet in _FA_SHEETS if sheet not in workbook.sheetnames]
            if missing:
                raise TranReferenceError(
                    "FA&GL workbook is missing required sheets: " + ", ".join(missing)
                )
            for sheet_name, (header_row, data_row, book, entity) in _FA_SHEETS.items():
                sheet = workbook[sheet_name]
                for column, expected in _FA_HEADERS.items():
                    actual = _normalized_text(sheet.cell(header_row, column).value)
                    if actual != _normalized_text(expected):
                        raise TranReferenceError(
                            f"{sheet_name}!{sheet.cell(header_row, column).coordinate} must be "
                            f"{expected!r}"
                        )
                for row_number, values in enumerate(
                    sheet.iter_rows(
                        min_row=data_row,
                        max_row=sheet.max_row,
                        min_col=1,
                        max_col=27,
                        values_only=True,
                    ),
                    start=data_row,
                ):
                    tag_number = _tag(values[15])
                    if not tag_number:
                        continue
                    indexed[tag_number].append(
                        FaGlRecord(
                            tag_number=tag_number,
                            asset_number=_plain_text(values[11]) or None,
                            start_date=_date_value(values[21]),
                            cost=_whole_vnd(values[26]),
                            book=book,
                            entity=entity,
                            cost_center=_plain_text(values[6]) or None,
                            product_code=_plain_text(values[7]) or None,
                            location=_plain_text(values[9]) or None,
                            company_name=_plain_text(values[1]) or None,
                            source_file=source_path.name,
                            source_sheet=sheet_name,
                            source_row=row_number,
                        )
                    )
        finally:
            workbook.close()
        return cls({key: tuple(value) for key, value in indexed.items()}, source_path)

    def lookup(self, tag_number: object) -> FaGlLookup:
        matches = self._records.get(_tag(tag_number), ())
        if not matches:
            status = ReferenceStatus.NOT_FOUND
        elif len(matches) == 1:
            status = ReferenceStatus.MATCHED
        else:
            status = ReferenceStatus.AMBIGUOUS
        return FaGlLookup(status=status, matches=matches)


@dataclass(frozen=True, slots=True)
class CcdcClassification:
    """Classification evidence read from Define or CMDB."""

    barcode: str
    product_type: str | None
    group_type: str | None
    group: DepreciationGroup | None
    physical: bool | None
    source_sheet: str
    source_row: int


@dataclass(frozen=True, slots=True)
class CcdcClassificationLookup:
    status: ReferenceStatus
    matches: tuple[CcdcClassification, ...]

    @property
    def classification(self) -> CcdcClassification | None:
        return self.matches[0] if self.status is ReferenceStatus.MATCHED else None


def _classification_group(
    barcode: str,
    product_type: object,
    group_type: object,
) -> tuple[DepreciationGroup | None, bool | None]:
    normalized_group = _normalized_text(group_type)
    if normalized_group in _NONPHYSICAL_GROUP_TYPES:
        return None, False
    if barcode in {"TPC", "CAM", "LEN", "PHO"}:
        return DepreciationGroup.SIX_YEAR, True
    if barcode in {"IPO", "SWA"}:
        return DepreciationGroup.FOUR_YEAR, True
    normalized_product = _normalized_text(product_type)
    if normalized_product in {"computer asset", "infrastructure asset"}:
        return DepreciationGroup.SIX_YEAR, True
    if normalized_product in {
        "computer component asset",
        "other",
        "spe part",
        "other spe part",
    }:
        return DepreciationGroup.FOUR_YEAR, True
    if normalized_group:
        return None, True
    return None, None


def _header_map(sheet: object, required: set[str]) -> tuple[int, dict[str, int]] | None:
    for row_number in range(1, min(getattr(sheet, "max_row", 1), 25) + 1):
        columns = {
            _normalized_text(sheet.cell(row_number, column).value): column
            for column in range(1, min(getattr(sheet, "max_column", 1), 100) + 1)
            if _plain_text(sheet.cell(row_number, column).value)
        }
        if required.issubset(columns):
            return row_number, columns
    return None


class CcdcWorkbookIndex:
    """Optional Define/CMDB/BC Xuatkho lookup without modifying the source workbook."""

    def __init__(
        self,
        classifications: dict[str, tuple[CcdcClassification, ...]],
        start_dates: dict[str, tuple[date, ...]],
        source_path: Path,
    ) -> None:
        self._classifications = classifications
        self._start_dates = start_dates
        self.source_path = source_path

    @classmethod
    def from_path(cls, path: str | Path) -> CcdcWorkbookIndex:
        source_path = Path(path).resolve()
        if not source_path.is_file() or source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise TranReferenceError("CCDC source must be an existing .xlsx or .xlsm file")
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        classifications: dict[str, list[CcdcClassification]] = defaultdict(list)
        start_dates: dict[str, list[date]] = defaultdict(list)
        try:
            if "Define" in workbook.sheetnames:
                sheet = workbook["Define"]
                found = _header_map(sheet, {"product type", "barcode", "group type"})
                if found is None:
                    raise TranReferenceError(
                        "Define must contain Product Type, Barcode, and Group Type headers"
                    )
                header_row, columns = found
                for row_number in range(header_row + 1, sheet.max_row + 1):
                    barcode = _tag(sheet.cell(row_number, columns["barcode"]).value)[:3]
                    if not barcode:
                        continue
                    product_type = _plain_text(
                        sheet.cell(row_number, columns["product type"]).value
                    )
                    group_type = _plain_text(
                        sheet.cell(row_number, columns["group type"]).value
                    )
                    group, physical = _classification_group(
                        barcode, product_type, group_type
                    )
                    classifications[barcode].append(
                        CcdcClassification(
                            barcode=barcode,
                            product_type=product_type or None,
                            group_type=group_type or None,
                            group=group,
                            physical=physical,
                            source_sheet="Define",
                            source_row=row_number,
                        )
                    )
            if "CMDB" in workbook.sheetnames:
                sheet = workbook["CMDB"]
                found = _header_map(sheet, {"asset name", "product type"})
                if found is None:
                    raise TranReferenceError("CMDB must contain Asset Name and Product Type")
                header_row, columns = found
                for row_number in range(header_row + 1, sheet.max_row + 1):
                    barcode = _tag(sheet.cell(row_number, columns["asset name"]).value)[:3]
                    if not barcode or barcode in classifications:
                        continue
                    product_type = _plain_text(
                        sheet.cell(row_number, columns["product type"]).value
                    )
                    group, physical = _classification_group(barcode, product_type, None)
                    classifications[barcode].append(
                        CcdcClassification(
                            barcode=barcode,
                            product_type=product_type or None,
                            group_type=None,
                            group=group,
                            physical=physical,
                            source_sheet="CMDB",
                            source_row=row_number,
                        )
                    )
            if "BC Xuatkho" in workbook.sheetnames:
                sheet = workbook["BC Xuatkho"]
                found = _header_map(sheet, {"asset name", "start time"})
                if found is None:
                    asset_column, start_column, header_row = 2, 5, 1
                else:
                    header_row, columns = found
                    asset_column = columns["asset name"]
                    start_column = columns["start time"]
                for row_number in range(header_row + 1, sheet.max_row + 1):
                    tag_number = _tag(sheet.cell(row_number, asset_column).value)
                    start_date = _date_value(sheet.cell(row_number, start_column).value)
                    if tag_number and start_date is not None:
                        start_dates[tag_number].append(start_date)
        finally:
            workbook.close()
        return cls(
            {key: tuple(value) for key, value in classifications.items()},
            {key: tuple(sorted(value)) for key, value in start_dates.items()},
            source_path,
        )

    def classification(self, barcode: object) -> CcdcClassificationLookup:
        matches = self._classifications.get(_tag(barcode)[:3], ())
        unique = {
            (item.product_type, item.group_type, item.group, item.physical) for item in matches
        }
        if not matches:
            status = ReferenceStatus.NOT_FOUND
        elif len(unique) == 1:
            status = ReferenceStatus.MATCHED
        else:
            status = ReferenceStatus.AMBIGUOUS
        return CcdcClassificationLookup(status=status, matches=matches)

    def earliest_start_date(self, tag_number: object) -> date | None:
        dates = self._start_dates.get(_tag(tag_number), ())
        return min(dates) if dates else None
