"""Safe accounting export that fills an existing import-workbook template."""

from __future__ import annotations

import os
import re
import shutil
import tempfile
from collections.abc import Iterable, Mapping
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from .accounting_xlsx import (
    AccountingExportError,
    AccountingValidationError,
    GlAccountPair,
    OutputExistsError,
)

ACCOUNTING_TEMPLATE_HEADERS = (
    "Invoice Type",
    "Invoice Number",
    "Supplier",
    "Supplier Site",
    "Invoice Date",
    "Line GL Date",
    "Invoice Amount",
    "Payment Method",
    "Head Description",
    "Line Description",
    "Line Num",
    "Line Amount",
    "Code Combination",
    "Tax Classification Code",
    "Included Tax Amount",
    "Supplier Name",
    "Supplier Tax",
    "Line Invoice Date FF",
    "Invoice Serial",
    "Invoice Num",
    "Item Description",
    "Valid Expense",
    "Batch Name",
    "Attribute Category",
    "Org Id",
    "Budget Code",
    "Reason Code",
    "ACCOUNT_FCT",
    "EFORM_NO",
    "Head GL Date",
)

_BUILT_IN_TEMPLATE = (
    Path(__file__).resolve().parents[1] / "templates" / "accounting_import_template.xlsx"
)
_DANGEROUS_EXCEL_PREFIXES = ("=", "+", "-", "@")
_GL_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/-]{0,127}$")


@dataclass(frozen=True, slots=True)
class AccountingTemplateExportResult:
    """Metadata for one completed template-based export."""

    path: Path
    template_path: Path
    output_suffix: str
    sheet_name: str
    batch_name: str
    invoice_date: date
    case_count: int
    invoice_count: int
    row_count: int
    source_total: Decimal
    credit_total: Decimal
    invoice_numbers: tuple[str, ...]

    @property
    def journal_line_count(self) -> int:
        """Compatibility alias for consumers that describe output rows as journal lines."""

        return self.row_count


@dataclass(frozen=True, slots=True)
class _PreparedCase:
    case_id: str
    case_type: str
    supplier_number: str | int
    supplier_site: str | int
    supplier_name: str
    description: str
    amount: Decimal
    debit_gl: str
    credit_lines: tuple[_CreditLine, ...]
    prepayment_highlight: str | None
    inactive: bool


@dataclass(frozen=True, slots=True)
class _CreditLine:
    account: str
    amount: Decimal
    highlight: str | None = None


@dataclass(frozen=True, slots=True)
class _RowArchetype:
    styles: tuple[object, ...]
    height: float | None
    hidden: bool
    outline_level: int
    collapsed: bool


def _case_value(case: object, name: str, default: Any = None) -> Any:
    if isinstance(case, Mapping):
        return case.get(name, default)
    return getattr(case, name, default)


def _enum_value(value: object) -> str:
    return str(getattr(value, "value", value) or "").strip().upper()


def _safe_text(value: object) -> str:
    text = str(value or "")
    if text.lstrip().startswith(_DANGEROUS_EXCEL_PREFIXES):
        return "'" + text
    return text


def _supplier_identifier(value: object) -> str | int:
    """Use numeric identifiers only when conversion cannot discard leading zeroes."""

    text = str(value or "").strip()
    if text.isdigit() and not text.startswith("0"):
        return int(text)
    return _safe_text(text)


def _whole_vnd(
    value: object,
    *,
    field_name: str,
    allow_none: bool = False,
) -> Decimal | None:
    if value in (None, ""):
        if allow_none:
            return None
        raise AccountingValidationError(f"{field_name} is required")
    if isinstance(value, bool):
        raise AccountingValidationError(f"{field_name} must be an exact whole VND amount")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise AccountingValidationError(f"Invalid {field_name}: {value!r}") from exc
    if not result.is_finite() or result != result.to_integral_value():
        raise AccountingValidationError(f"{field_name} must be an exact whole VND amount")
    if result < 0:
        raise AccountingValidationError(f"{field_name} cannot be negative")
    return result


def _gl_account(value: object, *, field_name: str) -> str:
    account = str(value or "").strip()
    if not _GL_RE.fullmatch(account):
        raise AccountingValidationError(f"Invalid {field_name}: {value!r}")
    return account


def _gl_pair(value: GlAccountPair | tuple[str, str]) -> GlAccountPair:
    if isinstance(value, GlAccountPair):
        debit_value, credit_value = value.debit, value.credit
    else:
        try:
            debit_value, credit_value = value
        except (TypeError, ValueError) as exc:
            raise AccountingValidationError(
                "GL mapping must contain debit/credit pairs"
            ) from exc
    debit = _gl_account(debit_value, field_name="debit GL")
    credit = _gl_account(credit_value, field_name="credit GL")
    if debit == credit:
        raise AccountingValidationError(f"Debit and credit GL cannot be the same: {debit}")
    return GlAccountPair(debit=debit, credit=credit)


def _credit_lines(
    raw_lines: object,
    *,
    amount: Decimal,
    fallback_credit: str,
    debit_gl: str,
) -> tuple[_CreditLine, ...]:
    if raw_lines in (None, ""):
        return (_CreditLine(fallback_credit, amount),)
    if not isinstance(raw_lines, (list, tuple)):
        raise AccountingValidationError(
            "credit_lines must be a list of account/amount pairs"
        )

    lines: list[_CreditLine] = []
    for index, raw_line in enumerate(raw_lines, start=1):
        if isinstance(raw_line, Mapping):
            raw_account = raw_line.get("account", raw_line.get("gl"))
            raw_amount = raw_line.get("amount")
            raw_highlight = raw_line.get("highlight")
        else:
            try:
                raw_account, raw_amount = raw_line
            except (TypeError, ValueError) as exc:
                raise AccountingValidationError(
                    f"Invalid credit line at position {index}"
                ) from exc
            raw_highlight = None
        account = _gl_account(
            raw_account,
            field_name=f"credit_lines[{index}].account",
        )
        if account == debit_gl:
            raise AccountingValidationError(
                f"Credit GL cannot equal debit GL at position {index}"
            )
        line_amount = _whole_vnd(
            raw_amount,
            field_name=f"credit_lines[{index}].amount",
        )
        assert line_amount is not None
        highlight = str(raw_highlight or "").strip().casefold() or None
        if highlight not in {None, "yellow", "green"}:
            raise AccountingValidationError(
                f"Invalid credit line highlight at position {index}"
            )
        lines.append(_CreditLine(account, line_amount, highlight))

    if not lines:
        raise AccountingValidationError("At least one credit line is required")
    line_total = sum((line.amount for line in lines), Decimal(0))
    if line_total != amount:
        raise AccountingValidationError("Credit lines do not reconcile to case amount")
    return tuple(lines)


def _invoice_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise AccountingValidationError("invoice_date must be a date or datetime")


def _capture_archetype(sheet: Worksheet, row: int) -> _RowArchetype:
    dimension = sheet.row_dimensions[row]
    cells = [sheet.cell(row, column) for column in range(1, len(ACCOUNTING_TEMPLATE_HEADERS) + 1)]
    return _RowArchetype(
        styles=tuple(copy(cell._style) for cell in cells),
        height=dimension.height,
        hidden=bool(dimension.hidden),
        outline_level=dimension.outlineLevel,
        collapsed=bool(dimension.collapsed),
    )


def _apply_archetype(sheet: Worksheet, row: int, archetype: _RowArchetype) -> None:
    source_dimension = archetype
    target_dimension = sheet.row_dimensions[row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outline_level
    target_dimension.collapsed = source_dimension.collapsed
    for column, style in enumerate(archetype.styles, start=1):
        target = sheet.cell(row, column)
        target._style = copy(style)
        target.value = None
        target.comment = None
        target.hyperlink = None


def _apply_highlight(sheet: Worksheet, row: int, highlight: str | None) -> None:
    colors = {"yellow": "FFFF00", "green": "92D050"}
    if highlight is None:
        return
    try:
        color = colors[highlight]
    except KeyError as exc:
        raise AccountingValidationError(f"Unsupported row highlight: {highlight}") from exc
    fill = PatternFill(fill_type="solid", fgColor=color)
    for column in range(1, len(ACCOUNTING_TEMPLATE_HEADERS) + 1):
        sheet.cell(row, column).fill = copy(fill)


def _atomic_commit(temp_path: Path, destination: Path, *, overwrite: bool) -> None:
    if overwrite:
        os.replace(temp_path, destination)
        return
    try:
        os.link(temp_path, destination)
    except FileExistsError as exc:
        raise OutputExistsError(f"Destination already exists: {destination}") from exc
    except OSError:
        created = False
        try:
            with destination.open("xb") as target, temp_path.open("rb") as source:
                created = True
                shutil.copyfileobj(source, target)
        except FileExistsError as exc:
            raise OutputExistsError(f"Destination already exists: {destination}") from exc
        except Exception:
            if created:
                destination.unlink(missing_ok=True)
            raise


class AccountingTemplateAdapter:
    """Fill a validated accounting-import template from Case-like objects.

    GL accounts are caller-supplied policy.  The adapter only maps cases into
    the template, retaining workbook structure and row archetype formatting.
    """

    def __init__(
        self,
        gl_mapping: Mapping[str, GlAccountPair | tuple[str, str]],
        *,
        template_path: str | Path | None = None,
        sheet_name: str | None = None,
        invoice_start: int = 1,
        invoice_width: int = 3,
        org_id: str | int | None = None,
    ) -> None:
        if not gl_mapping:
            raise AccountingValidationError("At least one GL mapping is required")
        self._gl_mapping = {
            _enum_value(case_type): _gl_pair(pair)
            for case_type, pair in gl_mapping.items()
        }
        resolved_template = Path(template_path) if template_path else _BUILT_IN_TEMPLATE
        if not resolved_template.is_file():
            raise FileNotFoundError(resolved_template)
        suffix = resolved_template.suffix.lower()
        if suffix not in {".xlsx", ".xlsm"}:
            raise AccountingValidationError("Template must be an .xlsx or .xlsm file")
        if (
            isinstance(invoice_start, bool)
            or not isinstance(invoice_start, int)
            or invoice_start < 0
        ):
            raise AccountingValidationError("invoice_start must be a nonnegative integer")
        if (
            isinstance(invoice_width, bool)
            or not isinstance(invoice_width, int)
            or invoice_width < 1
        ):
            raise AccountingValidationError("invoice_width must be a positive integer")
        self._template_path = resolved_template.resolve()
        self._sheet_name = sheet_name
        self._invoice_start = invoice_start
        self._invoice_width = invoice_width
        self._org_id = (
            None if org_id in (None, "") else _supplier_identifier(org_id)
        )

    @property
    def template_path(self) -> Path:
        return self._template_path

    @property
    def output_suffix(self) -> str:
        return self._template_path.suffix.lower()

    def _prepare_case(self, case: object) -> _PreparedCase:
        case_type = _enum_value(_case_value(case, "case_type"))
        if case_type not in {"DAMAGED", "LOST"}:
            raise AccountingValidationError(f"Unsupported case type: {case_type!r}")
        configured_gl = self._gl_mapping.get(case_type)
        if configured_gl is None:
            raise AccountingValidationError(f"Missing GL mapping for {case_type}")

        amount = _whole_vnd(_case_value(case, "amount"), field_name="amount")
        assert amount is not None
        residual_value = _whole_vnd(
            _case_value(case, "residual_value"),
            field_name="residual_value",
            allow_none=True,
        )
        responsibility_fee = _whole_vnd(
            _case_value(case, "responsibility_fee"),
            field_name="responsibility_fee",
            allow_none=True,
        )
        if (
            case_type == "LOST"
            and residual_value is not None
            and responsibility_fee is not None
            and residual_value + responsibility_fee != amount
        ):
            raise AccountingValidationError(
                "Lost-case residual value plus responsibility fee does not equal amount"
            )

        metadata_value = _case_value(case, "metadata", {}) or {}
        if not isinstance(metadata_value, Mapping):
            raise AccountingValidationError("case metadata must be a mapping")
        metadata = metadata_value
        debit_gl = configured_gl.debit
        fallback_credit = configured_gl.credit
        has_override = metadata.get("debit_gl") or metadata.get("credit_gl")
        if has_override:
            if not metadata.get("debit_gl") or not metadata.get("credit_gl"):
                raise AccountingValidationError(
                    "Both debit_gl and credit_gl overrides are required"
                )
            override = _gl_pair(
                (str(metadata["debit_gl"]), str(metadata["credit_gl"]))
            )
            debit_gl, fallback_credit = override.debit, override.credit

        raw_credit_lines = metadata.get(
            "credit_lines",
            _case_value(case, "credit_lines"),
        )
        credit_lines = _credit_lines(
            raw_credit_lines,
            amount=amount,
            fallback_credit=fallback_credit,
            debit_gl=debit_gl,
        )
        case_id = str(
            _case_value(case, "id")
            or _case_value(case, "source_id")
            or f"{case_type}:{_case_value(case, 'asset_code')}:{_case_value(case, 'domain')}"
        ).strip()
        if not case_id:
            raise AccountingValidationError("case id/source_id is required")

        description_value = metadata.get("description")
        if description_value in (None, ""):
            domain = str(_case_value(case, "domain") or "").strip()
            asset_code = str(_case_value(case, "asset_code") or "").strip()
            inactive_note = (
                " NV đã nghỉ việc"
                if metadata.get("employee_inactive") or metadata.get("inactive")
                else ""
            )
            if case_type == "LOST":
                description_value = (
                    f"Trừ lương {domain}{inactive_note} đền bù do thất lạc "
                    f"tài sản {asset_code}"
                )
            else:
                repair_status = _enum_value(_case_value(case, "repair_status"))
                repair_text = (
                    "có sửa chữa"
                    if repair_status in {"REPAIRED", "CÓ SỬA CHỮA", "CO SUA CHUA"}
                    else "không sửa chữa"
                    if repair_status
                    in {"NOT_REPAIRED", "KHÔNG SỬA CHỮA", "KHONG SUA CHUA"}
                    else "chưa xác định tình trạng sửa chữa"
                )
                description_value = (
                    f"Trừ lương {domain}{inactive_note} đền bù do hư hỏng "
                    f"tài sản {asset_code}, {repair_text}"
                )
        inactive = bool(metadata.get("employee_inactive") or metadata.get("inactive"))
        prepayment_highlight = (
            str(metadata.get("prepayment_highlight") or "").strip().casefold() or None
        )
        if prepayment_highlight not in {None, "yellow", "green"}:
            raise AccountingValidationError("Invalid prepayment_highlight")
        return _PreparedCase(
            case_id=case_id,
            case_type=case_type,
            supplier_number=_supplier_identifier(_case_value(case, "supplier_number")),
            supplier_site=_supplier_identifier(_case_value(case, "supplier_site")),
            supplier_name=_safe_text(_case_value(case, "supplier_name")),
            description=_safe_text(description_value),
            amount=amount,
            debit_gl=debit_gl,
            credit_lines=credit_lines,
            prepayment_highlight=prepayment_highlight,
            inactive=inactive,
        )

    def _select_sheet(self, workbook: Any) -> Worksheet:
        if self._sheet_name is None:
            return workbook.active
        if self._sheet_name not in workbook.sheetnames:
            raise AccountingValidationError(
                f"Template sheet not found: {self._sheet_name!r}"
            )
        return workbook[self._sheet_name]

    @staticmethod
    def _validate_headers(sheet: Worksheet) -> None:
        headers = tuple(
            sheet.cell(1, column).value
            for column in range(1, len(ACCOUNTING_TEMPLATE_HEADERS) + 1)
        )
        extra_headers = [
            sheet.cell(1, column).value
            for column in range(len(ACCOUNTING_TEMPLATE_HEADERS) + 1, sheet.max_column + 1)
            if sheet.cell(1, column).value not in (None, "")
        ]
        if headers != ACCOUNTING_TEMPLATE_HEADERS or extra_headers:
            raise AccountingValidationError(
                "Accounting template must contain the exact 30-column header contract"
            )

    @staticmethod
    def _clear_sample_data(sheet: Worksheet) -> int:
        original_max_row = max(sheet.max_row, 3)
        max_column = max(sheet.max_column, len(ACCOUNTING_TEMPLATE_HEADERS))
        for row in sheet.iter_rows(
            min_row=2,
            max_row=original_max_row,
            min_col=1,
            max_col=max_column,
        ):
            for cell in row:
                cell.value = None
                cell.comment = None
                cell.hyperlink = None
        return original_max_row

    @staticmethod
    def _write_mapped_row(
        sheet: Worksheet,
        row: int,
        values: Mapping[str, object],
    ) -> None:
        columns = {header: index for index, header in enumerate(ACCOUNTING_TEMPLATE_HEADERS, 1)}
        for header, value in values.items():
            sheet.cell(row, columns[header], value)

    @staticmethod
    def _resize_full_width_tables(sheet: Worksheet, last_data_row: int) -> None:
        for table in sheet.tables.values():
            min_column, min_row, max_column, _ = range_boundaries(table.ref)
            if (
                min_row == 1
                and min_column == 1
                and max_column == len(ACCOUNTING_TEMPLATE_HEADERS)
            ):
                table.ref = f"A1:{get_column_letter(max_column)}{last_data_row}"

    def _populate_sheet(
        self,
        sheet: Worksheet,
        cases: list[_PreparedCase],
        *,
        batch_name: str,
        invoice_date: date,
    ) -> tuple[tuple[str, ...], int, int]:
        self._validate_headers(sheet)
        prepayment = _capture_archetype(sheet, 2)
        credit = _capture_archetype(sheet, 3)
        prepayment_type = str(sheet.cell(2, 1).value or "").strip()
        credit_type = str(sheet.cell(3, 1).value or "").strip()
        if not prepayment_type or not credit_type:
            raise AccountingValidationError(
                "Template rows 2 and 3 must define prepayment and credit invoice types"
            )
        original_max_row = self._clear_sample_data(sheet)

        row = 2
        invoice_sequence = self._invoice_start
        invoice_numbers: list[str] = []
        for case in cases:
            prepayment_number = f"{batch_name}{invoice_sequence:0{self._invoice_width}d}"
            invoice_sequence += 1
            credit_number = f"{batch_name}{invoice_sequence:0{self._invoice_width}d}"
            invoice_sequence += 1
            invoice_numbers.extend((prepayment_number, credit_number))

            _apply_archetype(sheet, row, prepayment)
            _apply_highlight(
                sheet,
                row,
                case.prepayment_highlight or ("yellow" if case.inactive else None),
            )
            self._write_mapped_row(
                sheet,
                row,
                {
                    "Invoice Type": _safe_text(prepayment_type),
                    "Invoice Number": _safe_text(prepayment_number),
                    "Supplier": case.supplier_number,
                    "Supplier Site": case.supplier_site,
                    "Invoice Date": invoice_date,
                    "Line GL Date": invoice_date,
                    "Invoice Amount": int(case.amount),
                    "Head Description": case.description,
                    "Line Description": case.description,
                    "Line Num": 1,
                    "Line Amount": int(case.amount),
                    "Code Combination": _safe_text(case.debit_gl),
                    "Supplier Name": case.supplier_name,
                    "Batch Name": _safe_text(batch_name),
                    "Org Id": self._org_id,
                },
            )
            row += 1

            for line_number, credit_line in enumerate(
                case.credit_lines,
                start=1,
            ):
                _apply_archetype(sheet, row, credit)
                _apply_highlight(
                    sheet,
                    row,
                    credit_line.highlight or ("yellow" if case.inactive else None),
                )
                self._write_mapped_row(
                    sheet,
                    row,
                    {
                        "Invoice Type": _safe_text(credit_type),
                        "Invoice Number": _safe_text(credit_number),
                        "Supplier": case.supplier_number,
                        "Supplier Site": case.supplier_site,
                        "Invoice Date": invoice_date,
                        "Line GL Date": invoice_date,
                        "Invoice Amount": -int(case.amount),
                        "Head Description": case.description,
                        "Line Description": case.description,
                        "Line Num": line_number,
                        "Line Amount": -int(credit_line.amount),
                        "Code Combination": _safe_text(credit_line.account),
                        "Supplier Name": case.supplier_name,
                        "Batch Name": _safe_text(batch_name),
                        "Org Id": self._org_id,
                    },
                )
                row += 1

        last_data_row = row - 1
        for stale_row in range(last_data_row + 1, original_max_row + 1):
            for column in range(1, len(ACCOUNTING_TEMPLATE_HEADERS) + 1):
                sheet.cell(stale_row, column).value = None
        self._resize_full_width_tables(sheet, last_data_row)
        return tuple(invoice_numbers), last_data_row - 1, last_data_row

    def _destination_path(self, destination: str | Path) -> Path:
        path = Path(destination)
        if not path.suffix:
            return path.with_suffix(self.output_suffix)
        if path.suffix.lower() != self.output_suffix:
            raise AccountingValidationError(
                f"Destination extension must match template: {self.output_suffix}"
            )
        return path

    def export(
        self,
        cases: Iterable[object],
        destination: str | Path,
        *,
        batch_name: str,
        invoice_date: date | datetime,
        expected_total: Decimal | int | None = None,
        overwrite: bool = False,
    ) -> AccountingTemplateExportResult:
        """Export two invoices per case, with one row per configured credit line."""

        prepared = [self._prepare_case(case) for case in cases]
        if not prepared:
            raise AccountingValidationError("Cannot export an empty accounting batch")
        case_ids = [case.case_id for case in prepared]
        duplicate_ids = sorted(
            {case_id for case_id in case_ids if case_ids.count(case_id) > 1}
        )
        if duplicate_ids:
            raise AccountingValidationError(
                f"Duplicate case ids: {', '.join(duplicate_ids)}"
            )
        raw_batch_name = str(batch_name or "").strip()
        if not raw_batch_name:
            raise AccountingValidationError("batch_name is required")
        export_date = _invoice_date(invoice_date)
        source_total = sum((case.amount for case in prepared), Decimal(0))
        expected = (
            source_total
            if expected_total is None
            else _whole_vnd(expected_total, field_name="expected_total")
        )
        assert expected is not None
        if expected != source_total:
            raise AccountingValidationError(
                f"Source total {source_total} does not match expected total {expected}"
            )
        credit_total = sum(
            (line.amount for case in prepared for line in case.credit_lines),
            Decimal(0),
        )
        if credit_total != source_total:
            raise AccountingValidationError("Credit lines do not reconcile to source total")

        destination_path = self._destination_path(destination)
        if destination_path.exists() and not overwrite:
            raise OutputExistsError(f"Destination already exists: {destination_path}")
        destination_path.parent.mkdir(parents=True, exist_ok=True)
        keep_vba = self.output_suffix == ".xlsm"
        workbook = load_workbook(self.template_path, keep_vba=keep_vba)
        try:
            sheet = self._select_sheet(workbook)
            invoice_numbers, row_count, last_data_row = self._populate_sheet(
                sheet,
                prepared,
                batch_name=raw_batch_name,
                invoice_date=export_date,
            )
            with tempfile.NamedTemporaryFile(
                prefix=f".{destination_path.stem}.",
                suffix=self.output_suffix,
                dir=destination_path.parent,
                delete=False,
            ) as temp_handle:
                temp_path = Path(temp_handle.name)
            try:
                workbook.save(temp_path)
                self._verify_output(
                    temp_path,
                    sheet_name=sheet.title,
                    expected_row_count=row_count,
                    last_data_row=last_data_row,
                    keep_vba=keep_vba,
                )
                _atomic_commit(temp_path, destination_path, overwrite=overwrite)
            finally:
                temp_path.unlink(missing_ok=True)
        finally:
            workbook.close()

        return AccountingTemplateExportResult(
            path=destination_path,
            template_path=self.template_path,
            output_suffix=self.output_suffix,
            sheet_name=sheet.title,
            batch_name=raw_batch_name,
            invoice_date=export_date,
            case_count=len(prepared),
            invoice_count=len(prepared) * 2,
            row_count=row_count,
            source_total=source_total,
            credit_total=credit_total,
            invoice_numbers=invoice_numbers,
        )

    @classmethod
    def _verify_output(
        cls,
        path: Path,
        *,
        sheet_name: str,
        expected_row_count: int,
        last_data_row: int,
        keep_vba: bool,
    ) -> None:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
        try:
            if sheet_name not in workbook.sheetnames:
                raise AccountingExportError("Saved workbook is missing the target sheet")
            sheet = workbook[sheet_name]
            cls._validate_headers(sheet)
            invoice_column = ACCOUNTING_TEMPLATE_HEADERS.index("Invoice Number") + 1
            populated_rows = sum(
                sheet.cell(row, invoice_column).value not in (None, "")
                for row in range(2, last_data_row + 1)
            )
            if populated_rows != expected_row_count:
                raise AccountingExportError("Saved workbook row count does not reconcile")
        finally:
            workbook.close()
