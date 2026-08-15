"""No-clobber exporter for the external TranNNB calculation template."""

from __future__ import annotations

import os
import shutil
import tempfile
from collections.abc import Iterable
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.worksheet import Worksheet

from asset_compensation.domain import CompensationStatus, DepreciationGroup, ValidationError
from asset_compensation.services.compensation_service import FOUR_YEAR_RATES, SIX_YEAR_RATES

from .accounting_xlsx import OutputExistsError

if TYPE_CHECKING:
    from asset_compensation.services.tran_workflow_service import TranResolution

TRAN_YEAR_HEADERS = (
    "No.",
    "Tháng đền bù",
    "Asset Number",
    "Tagnumber",
    "Tên tài sản",
    "Nhânviênđềnbù",
    "Ngày đưa vào sử dụng",
    "Ngày mất",
    "Nguyên giá ban đầu (vnđ)",
    "Giá trị còn lại (vnđ)",
    "Phí đền bù trách nhiệm (vnd)",
    "Tổng tiền nhân viên phải hoàn trả cho công ty (vnđ)",
    "Thời gian đã sử dụng (tháng)",
    "Sổ",
    "Entity",
    "Cost center",
    "Product code",
    "Location",
    "Thời gian sử dụng còn lại (tháng)",
    "Tỷ lệ \nchi phí đền bù (%)",
    "ORC",
    "Note",
)
TRAN_SENT_HEADERS = TRAN_YEAR_HEADERS[3:18]
_DANGEROUS_EXCEL_PREFIXES = ("=", "+", "-", "@")


class TranWorkbookError(ValidationError):
    """Raised when a Tran template or export request is not safe to process."""


@dataclass(frozen=True, slots=True)
class TranWorkbookExportResult:
    path: Path
    template_path: Path
    year_sheet: str
    request_rows: tuple[int, ...]
    sent_out_rows: int


def _safe_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value)
    if text.lstrip().startswith(_DANGEROUS_EXCEL_PREFIXES):
        return "'" + text
    return text


def _percent(rate: object) -> str:
    return f"{int(float(rate) * 100)}%"


def remaining_value_excel_formula(
    row: int,
    group: DepreciationGroup,
    usage_months: int,
) -> str:
    """Build the auditable full-year formula required by the original workflow."""

    current_year = usage_months // 12 + 1
    remaining_months = 12 - usage_months % 12
    rates = SIX_YEAR_RATES if group is DepreciationGroup.SIX_YEAR else FOUR_YEAR_RATES
    terminal_year = 6 if group is DepreciationGroup.SIX_YEAR else 4
    if current_year > terminal_year:
        return f"=ROUND(I{row}*10%,0)"
    terms: list[str] = []
    for year, rate in enumerate(rates, start=1):
        label = _percent(rate)
        if year < current_year:
            terms.append(f"{label}*(0/12)")
        elif year == current_year:
            terms.append(f"{label}*({remaining_months}/12)")
        else:
            terms.append(label)
    if group is DepreciationGroup.SIX_YEAR or current_year == 4:
        terms.append("10%")
    return f"=ROUND(I{row}*({' + '.join(terms)}),0)"


def _atomic_commit(temp_path: Path, destination: Path) -> None:
    try:
        os.link(temp_path, destination)
    except FileExistsError as exc:
        raise OutputExistsError(f"Destination already exists: {destination}") from exc
    except OSError:
        created = False
        try:
            with destination.open("xb") as target, temp_path.open("rb") as source:
                created = True
                shutil.copyfileobj(source, target)
        except FileExistsError as exc:
            raise OutputExistsError(f"Destination already exists: {destination}") from exc
        except Exception:
            if created:
                destination.unlink(missing_ok=True)
            raise


def _copy_row_style(sheet: Worksheet, source_row: int, target_row: int) -> None:
    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    for column in range(1, len(TRAN_YEAR_HEADERS) + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.value = None
        target.comment = None
        target.hyperlink = None


def _last_request_row(sheet: Worksheet) -> int:
    for row in range(sheet.max_row, 3, -1):
        if sheet.cell(row, 4).value not in (None, ""):
            return row
    return 3


def _next_number(sheet: Worksheet, last_row: int) -> int:
    for row in range(last_row, 3, -1):
        value = sheet.cell(row, 1).value
        if isinstance(value, int):
            return value + 1
        if isinstance(value, float) and value.is_integer():
            return int(value) + 1
    return 1


def _remaining_usage_months(group: DepreciationGroup | None, used: int | None) -> int | None:
    if group is None or used is None:
        return None
    schedule = 72 if group is DepreciationGroup.SIX_YEAR else 48
    return max(schedule - used, 0)


class TranWorkbookAdapter:
    """Append one request to a year sheet and rebuild request-only ``Sent out``."""

    def __init__(
        self,
        template_path: str | Path,
        *,
        sent_out_sheet: str = "Sent out",
    ) -> None:
        path = Path(template_path).resolve()
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise TranWorkbookError("Tran template must be an existing .xlsx or .xlsm file")
        self.template_path = path
        self.sent_out_sheet = sent_out_sheet

    @property
    def output_suffix(self) -> str:
        return self.template_path.suffix.lower()

    @staticmethod
    def _validate_year_sheet(sheet: Worksheet) -> None:
        headers = tuple(sheet.cell(3, column).value for column in range(1, 23))
        if headers != TRAN_YEAR_HEADERS:
            raise TranWorkbookError("Tran year sheet does not match the exact A:V header contract")

    @staticmethod
    def _validate_resolutions(
        resolutions: Iterable[TranResolution],
    ) -> list[TranResolution]:
        result = list(resolutions)
        if not result:
            raise TranWorkbookError("Cannot export an empty Tran request")
        for index, resolution in enumerate(result, start=1):
            if resolution.asset is None or resolution.preview is None:
                raise TranWorkbookError(f"Tran item {index} has unresolved reference fields")
            if resolution.preview.status is CompensationStatus.NEEDS_REVIEW:
                raise TranWorkbookError(f"Tran item {index} still requires manual review")
        return result

    @staticmethod
    def _write_year_row(
        sheet: Worksheet,
        row: int,
        number: int,
        resolution: TranResolution,
        processing_date: date,
    ) -> None:
        asset = resolution.asset
        preview = resolution.preview
        assert asset is not None and preview is not None
        values = {
            1: number,
            2: date(processing_date.year, processing_date.month, 1),
            3: _safe_text(asset.asset_number),
            4: _safe_text(asset.tag_number),
            5: _safe_text(asset.asset_name),
            6: _safe_text(asset.domain),
            7: asset.start_date,
            8: asset.lost_date,
            9: int(asset.cost) if asset.cost is not None else None,
            13: f"=ROUND(DAYS360(G{row},H{row},TRUE)/30,0)",
            14: _safe_text(asset.book),
            15: _safe_text(asset.entity),
            16: _safe_text(asset.cost_center),
            17: _safe_text(asset.product_code),
            18: _safe_text(asset.location),
            19: _remaining_usage_months(preview.depreciation_group, preview.usage_months),
            20: float(preview.fee_rate) if preview.fee_rate is not None else None,
            21: None,
            22: _safe_text(" ".join(resolution.notes)),
        }
        if preview.status is CompensationStatus.CALCULATED:
            assert preview.depreciation_group is not None and preview.usage_months is not None
            values[10] = remaining_value_excel_formula(
                row, preview.depreciation_group, preview.usage_months
            )
            values[11] = f"=ROUND(I{row}*T{row},0)"
            values[12] = f"=ROUND(SUM(J{row}:K{row}),0)"
        elif preview.status is CompensationStatus.EXEMPT:
            values[10] = "Không tính đền bù"
            values[11] = None
            values[12] = None
        elif preview.status is CompensationStatus.NOT_APPLICABLE:
            values[10] = "Không áp dụng"
            values[11] = None
            values[12] = None
        for column, value in values.items():
            sheet.cell(row, column, value)

    @staticmethod
    def _sent_values(resolution: TranResolution) -> tuple[object, ...]:
        asset = resolution.asset
        preview = resolution.preview
        assert asset is not None and preview is not None
        remaining: object = preview.remaining_value
        fee_value: object = preview.fee_value
        total_amount: object = preview.total_amount
        if preview.status is CompensationStatus.EXEMPT:
            remaining = "Không tính đền bù"
            fee_value = None
            total_amount = None
        elif preview.status is CompensationStatus.NOT_APPLICABLE:
            remaining = "Không áp dụng"
            fee_value = None
            total_amount = None
        return (
            _safe_text(asset.tag_number),
            _safe_text(asset.asset_name),
            _safe_text(asset.domain),
            asset.start_date,
            asset.lost_date,
            int(asset.cost) if asset.cost is not None else None,
            remaining,
            fee_value,
            total_amount,
            preview.usage_months,
            _safe_text(asset.book),
            _safe_text(asset.entity),
            _safe_text(asset.cost_center),
            _safe_text(asset.product_code),
            _safe_text(asset.location),
        )

    @staticmethod
    def _rebuild_sent_out(sheet: Worksheet, resolutions: list[TranResolution]) -> None:
        for merged_range in tuple(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged_range))
        if sheet.max_row:
            sheet.delete_rows(1, sheet.max_row)
        for column, header in enumerate(TRAN_SENT_HEADERS, start=1):
            sheet.cell(1, column, header)
        for row_number, resolution in enumerate(resolutions, start=2):
            for column, value in enumerate(
                TranWorkbookAdapter._sent_values(resolution), start=1
            ):
                sheet.cell(row_number, column, value)

        last_data_row = len(resolutions) + 1
        total_row = last_data_row + 1
        sheet.cell(total_row, 2, "Total:")
        sheet.cell(total_row, 7, f"=SUM(G2:G{last_data_row})")
        sheet.cell(total_row, 8, f"=SUM(H2:H{last_data_row})")
        sheet.cell(total_row, 9, f"=SUM(I2:I{last_data_row})")

        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        light_blue = PatternFill(
            fill_type="solid", fgColor=Color(theme=4, tint=0.7999)
        )
        for cell in sheet[1][:15]:
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.fill = light_blue
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = border
        for row in range(2, total_row):
            for column in range(1, 16):
                cell = sheet.cell(row, column)
                cell.font = Font(name="Arial", size=10)
                cell.border = border
                cell.alignment = Alignment(vertical="center")
            for column in (4, 5, 6, 7, 8, 9):
                sheet.cell(row, column).alignment = Alignment(
                    horizontal="right", vertical="center"
                )
            for column in (10, 11, 12, 13, 14, 15):
                sheet.cell(row, column).alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            for column in (4, 5):
                sheet.cell(row, column).number_format = "dd/mm/yyyy"
            for column in (6, 7, 8, 9):
                sheet.cell(row, column).number_format = "#,##0"
        for column in range(1, 16):
            cell = sheet.cell(total_row, column)
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.fill = light_blue
            cell.border = border
        for column in (7, 8, 9):
            sheet.cell(total_row, column).number_format = "#,##0"
            sheet.cell(total_row, column).alignment = Alignment(
                horizontal="right", vertical="center"
            )

    def export(
        self,
        resolutions: Iterable[TranResolution],
        destination: str | Path,
        *,
        processing_date: date,
        year_sheet: str | None = None,
    ) -> TranWorkbookExportResult:
        prepared = self._validate_resolutions(resolutions)
        destination_path = Path(destination).resolve()
        if destination_path == self.template_path:
            raise TranWorkbookError("The source Tran template can never be overwritten")
        if destination_path.suffix.lower() != self.output_suffix:
            raise TranWorkbookError(
                f"Destination must preserve the template suffix {self.output_suffix}"
            )
        if destination_path.exists():
            raise OutputExistsError(f"Destination already exists: {destination_path}")
        destination_path.parent.mkdir(parents=True, exist_ok=True)

        sheet_name = year_sheet or str(processing_date.year)
        workbook = load_workbook(
            self.template_path,
            data_only=False,
            keep_vba=self.output_suffix == ".xlsm",
        )
        try:
            if sheet_name not in workbook.sheetnames:
                raise TranWorkbookError(f"Tran year sheet not found: {sheet_name}")
            if self.sent_out_sheet not in workbook.sheetnames:
                raise TranWorkbookError(
                    f"Tran Sent out sheet not found: {self.sent_out_sheet}"
                )
            year = workbook[sheet_name]
            sent_out = workbook[self.sent_out_sheet]
            self._validate_year_sheet(year)
            last_row = _last_request_row(year)
            first_row = last_row + 1
            number = _next_number(year, last_row)
            source_style_row = 4
            request_rows: list[int] = []
            for offset, resolution in enumerate(prepared):
                row = first_row + offset
                _copy_row_style(year, source_style_row, row)
                self._write_year_row(year, row, number + offset, resolution, processing_date)
                request_rows.append(row)
            self._rebuild_sent_out(sent_out, prepared)
            if workbook.calculation is None:
                workbook.calculation = CalcProperties()
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = "auto"
            with tempfile.NamedTemporaryFile(
                prefix=f".{destination_path.stem}.",
                suffix=self.output_suffix,
                dir=destination_path.parent,
                delete=False,
            ) as temp_handle:
                temp_path = Path(temp_handle.name)
            try:
                workbook.save(temp_path)
                self._verify(temp_path, sheet_name, tuple(request_rows), len(prepared))
                _atomic_commit(temp_path, destination_path)
            finally:
                temp_path.unlink(missing_ok=True)
        finally:
            workbook.close()

        return TranWorkbookExportResult(
            path=destination_path,
            template_path=self.template_path,
            year_sheet=sheet_name,
            request_rows=tuple(request_rows),
            sent_out_rows=len(prepared),
        )

    def _verify(
        self,
        path: Path,
        year_sheet: str,
        request_rows: tuple[int, ...],
        sent_count: int,
    ) -> None:
        workbook = load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_vba=path.suffix.lower() == ".xlsm",
        )
        try:
            year = workbook[year_sheet]
            self._validate_year_sheet(year)
            sent = workbook[self.sent_out_sheet]
            headers = tuple(sent.cell(1, column).value for column in range(1, 16))
            if headers != TRAN_SENT_HEADERS:
                raise TranWorkbookError("Saved Sent out header contract is invalid")
            if sent.max_row != sent_count + 2:
                raise TranWorkbookError("Saved Sent out contains stale or blank rows")
            for row in request_rows:
                if year.cell(row, 4).value in (None, ""):
                    raise TranWorkbookError("Saved year sheet is missing a request row")
        finally:
            workbook.close()
