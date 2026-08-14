"""Strict supplier-directory loaders for CSV and modern Excel files."""

from __future__ import annotations

import csv
import unicodedata
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .contracts import SupplierRecord


class SupplierLoadError(ValueError):
    """Raised when supplier source data does not satisfy the contract."""


class DuplicateSupplierDomainError(SupplierLoadError):
    """Raised instead of silently applying last-write-wins semantics."""

    def __init__(self, domain: str, first_row: int | None, duplicate_row: int | None) -> None:
        self.domain = domain
        self.first_row = first_row
        self.duplicate_row = duplicate_row
        super().__init__(
            f"Duplicate supplier domain {domain!r} at rows {first_row} and {duplicate_row}"
        )


_ALIASES = {
    "domain": {"domain", "user", "username", "email", "user domain"},
    "supplier_number": {
        "supplier number",
        "supplier no",
        "supplier code",
        "vendor number",
        "vendor code",
        "ma nha cung cap",
    },
    "supplier_site": {"supplier site", "vendor site", "site", "dia diem nha cung cap"},
    "supplier_name": {"supplier name", "vendor name", "name", "ten nha cung cap"},
    "active": {"active", "is active", "status", "trang thai"},
}


def normalize_domain(value: object) -> str:
    domain = str(value or "").strip().casefold()
    if "@" in domain:
        domain = domain.partition("@")[0]
    return domain


def _fold_header(value: object) -> str:
    text = str(value or "").strip().casefold().replace("đ", "d")
    normalized = unicodedata.normalize("NFKD", text)
    return " ".join(
        "".join(char for char in normalized if not unicodedata.combining(char)).split()
    )


def _canonical_headers(headers: Iterable[object]) -> dict[str, int]:
    folded = [_fold_header(value) for value in headers]
    result: dict[str, int] = {}
    for canonical, aliases in _ALIASES.items():
        normalized_aliases = {_fold_header(alias) for alias in aliases}
        matches = [index for index, header in enumerate(folded) if header in normalized_aliases]
        if len(matches) > 1:
            raise SupplierLoadError(f"Ambiguous columns for {canonical!r}: {matches}")
        if matches:
            result[canonical] = matches[0]
    missing = {"domain", "supplier_number", "supplier_site"} - result.keys()
    if missing:
        raise SupplierLoadError(f"Missing supplier columns: {', '.join(sorted(missing))}")
    return result


def _as_bool(value: object) -> bool | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    folded = _fold_header(value)
    if folded in {"1", "true", "yes", "y", "active", "hoat dong"}:
        return True
    if folded in {"0", "false", "no", "n", "inactive", "khong hoat dong"}:
        return False
    raise SupplierLoadError(f"Unsupported active value: {value!r}")


def _row_value(row: list[object], columns: Mapping[str, int], name: str) -> object:
    index = columns.get(name)
    return row[index] if index is not None and index < len(row) else None


def _records_from_rows(rows: Iterable[Iterable[object]]) -> list[SupplierRecord]:
    iterator = iter(rows)
    try:
        headers = list(next(iterator))
    except StopIteration as exc:
        raise SupplierLoadError("Supplier source is empty") from exc
    columns = _canonical_headers(headers)
    records: list[SupplierRecord] = []
    seen: dict[str, int] = {}
    for row_number, raw_row in enumerate(iterator, start=2):
        row = list(raw_row)
        if not any(value not in (None, "") for value in row):
            continue

        domain = normalize_domain(_row_value(row, columns, "domain"))
        if not domain:
            raise SupplierLoadError(f"Missing domain at row {row_number}")
        if domain in seen:
            raise DuplicateSupplierDomainError(domain, seen[domain], row_number)
        seen[domain] = row_number

        supplier_number = str(_row_value(row, columns, "supplier_number") or "").strip()
        supplier_number = supplier_number.lstrip("'")
        supplier_site = str(_row_value(row, columns, "supplier_site") or "").strip()
        supplier_site = supplier_site.lstrip("'")
        if not supplier_number or not supplier_site:
            raise SupplierLoadError(f"Missing supplier number/site at row {row_number}")

        supplier_name_value = _row_value(row, columns, "supplier_name")
        records.append(
            SupplierRecord(
                domain=domain,
                supplier_number=supplier_number,
                supplier_site=supplier_site,
                supplier_name=(
                    str(supplier_name_value).strip()
                    if supplier_name_value not in (None, "")
                    else None
                ),
                active=_as_bool(_row_value(row, columns, "active")),
                source_row=row_number,
            )
        )
    return records


def _mapping_rows(rows: Iterable[Mapping[str, Any]]) -> Iterable[list[object]]:
    materialized = list(rows)
    if not materialized:
        return []
    headers = list(materialized[0].keys())
    return [headers, *[[row.get(header) for header in headers] for row in materialized]]


def load_supplier_records(
    source: str | Path | Iterable[Mapping[str, Any]],
) -> list[SupplierRecord]:
    """Load supplier rows and fail fast on duplicate normalized domains."""

    if not isinstance(source, (str, Path)):
        return _records_from_rows(_mapping_rows(source))

    path = Path(source)
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return _records_from_rows(csv.reader(stream))
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            return _records_from_rows(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    if suffix == ".xls":
        raise SupplierLoadError(
            "Legacy .xls is not supported by the safe core loader; convert it to .xlsx or .csv"
        )
    raise SupplierLoadError(f"Unsupported supplier source: {path.suffix or '<no extension>'}")


def load_supplier_directory(
    source: str | Path | Iterable[Mapping[str, Any]],
) -> dict[str, SupplierRecord]:
    records = load_supplier_records(source)
    return {record.domain: record for record in records}
