"""HTTP contracts for Outlook add-in and local bridge companions."""

from __future__ import annotations

import base64
import io
from datetime import date
from email.message import EmailMessage
from pathlib import Path

from flask import Flask
from openpyxl import Workbook

from asset_compensation.config import Settings
from asset_compensation.web import create_app

_BASIC = {
    "Authorization": "Basic "
    + base64.b64encode(b"demo-team:synthetic-secret").decode("ascii")
}


def _fa_gl_bytes() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {
        "VNG-Asset": (3, 4),
        "VNG-Tool": (3, 4),
        "VNGS-Asset": (2, 3),
        "VNGS-Tool": (2, 3),
    }
    headers = {
        2: "Company Name",
        7: "Cost Center",
        8: "Product Code",
        10: "Location",
        12: "Asset",
        16: "Tag Number",
        22: "Date of Depreciation",
        25: "Life(month)",
        27: "Cost",
    }
    for sheet_name, (header_row, data_row) in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for column, header in headers.items():
            sheet.cell(header_row, column, header)
        if sheet_name == "VNG-Tool":
            sheet.cell(data_row, 2, "Synthetic Company")
            sheet.cell(data_row, 7, "0603")
            sheet.cell(data_row, 8, "000")
            sheet.cell(data_row, 10, "01")
            sheet.cell(data_row, 12, "SYN-001")
            sheet.cell(data_row, 16, "LAP10001")
            sheet.cell(data_row, 22, date(2025, 1, 1))
            sheet.cell(data_row, 25, 48)
            sheet.cell(data_row, 27, 2_000_000)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _lost_email() -> bytes:
    message = EmailMessage()
    message["Subject"] = "IT - Thông tin tài sản thất lạc"
    message["From"] = "Synthetic Asset Team <asset@example.invalid>"
    message["To"] = "Operator <operator@example.invalid>"
    message["Cc"] = "Audit <audit@example.invalid>"
    message["Date"] = "Thu, 13 Aug 2026 03:49:56 +0000"
    message["Message-ID"] = "<synthetic-companion@example.invalid>"
    message.set_content("Synthetic lost-asset table follows.")
    message.add_alternative(
        """
        <html><body><table>
          <tr>
            <th>Asset Name</th><th>Product Name</th><th>Domain</th>
            <th>Ngày bắt đầu sử dụng</th><th>Ngày thất lạc/mất</th>
            <th>Nguyên giá ban đầu</th><th>Mức khấu hao sử dụng còn lại</th>
            <th>Phí đền bù trách nhiệm</th><th>Tổng số tiền đền bù</th>
            <th>NOTE</th><th>Entity</th><th>Cost center</th>
            <th>Product code</th><th>Location</th>
          </tr>
          <tr>
            <td>LAP10001</td><td>Synthetic Laptop</td><td>demo.alpha</td>
            <td>01/01/2025</td><td>01/08/2026</td><td>2,000,000</td>
            <td>600,000</td><td>100,000</td><td>700,000</td><td>Asset</td><td>VNG</td>
            <td>'0603</td><td>'000</td><td>'01</td>
          </tr>
        </table></body></html>
        """,
        subtype="html",
    )
    return message.as_bytes()


def _app(tmp_path: Path) -> Flask:
    fa_path = tmp_path / "synthetic-fa.xlsx"
    fa_path.write_bytes(_fa_gl_bytes())
    app = create_app(
        Settings(
            data_dir=tmp_path / "runtime",
            demo_mode=False,
            allow_test_reset=True,
            retain_raw_eml=True,
            fa_gl_reference=fa_path,
            access_user="demo-team",
            access_password="synthetic-secret",
            secret_key="synthetic-companion-key",
        )
    )
    app.config.update(TESTING=True)
    return app


def _pair(client: object, role: str, client_type: str = "local_bridge") -> dict[str, object]:
    paired = client.post(  # type: ignore[attr-defined]
        "/api/companion/pairings",
        json={"role": role, "client_type": client_type},
        headers={**_BASIC, "X-Asset-Hub-Action": "companion-pair-v1"},
    )
    assert paired.status_code == 201
    code = paired.get_json()["pairing"]["code"]
    exchanged = client.post(  # type: ignore[attr-defined]
        "/api/companion/exchange",
        json={"code": code},
    )
    assert exchanged.status_code == 200
    return exchanged.get_json()


def test_pairing_endpoints_have_narrow_basic_and_bearer_boundaries(tmp_path: Path) -> None:
    app = _app(tmp_path)
    try:
        client = app.test_client()
        denied = client.post(
            "/api/companion/pairings",
            json={"role": "tran", "client_type": "outlook_addin"},
            headers={"X-Asset-Hub-Action": "companion-pair-v1"},
        )
        assert denied.status_code == 401
        assert denied.headers["WWW-Authenticate"].startswith("Basic")

        invalid_exchange = client.post("/api/companion/exchange", json={"code": "invalid"})
        assert invalid_exchange.status_code == 401
        assert invalid_exchange.headers["WWW-Authenticate"].startswith("Bearer")
        assert invalid_exchange.headers["Cache-Control"] == "private, no-store"

        exchange = _pair(client, "tran", "outlook_addin")
        assert exchange["role"] == "tran"
        assert exchange["client_type"] == "outlook_addin"
        assert exchange["expires_in_seconds"] == 28_800
        token = exchange["token"]
        assert isinstance(token, str)

        missing_bearer = client.post(
            "/api/companion/client/emails",
            data=b"not parsed",
            content_type="application/octet-stream",
        )
        assert missing_bearer.status_code == 401
        assert missing_bearer.headers["WWW-Authenticate"].startswith("Bearer")

        future_client_route = client.get(
            "/api/companion/client/future-route",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert future_client_route.status_code == 401
        assert future_client_route.headers["WWW-Authenticate"].startswith("Basic")

        capabilities = client.get("/api/capabilities", headers=_BASIC).get_json()[
            "capabilities"
        ]
        assert capabilities["companion_pairing"] is True
        assert capabilities["outlook_addin"] is True
        assert capabilities["local_bridge"] is True
        assert capabilities["tran_companion_draft"] is True
    finally:
        app.extensions["asset_hub"]["repository"].close()


def test_tran_companion_upload_draft_queue_and_reset_are_source_bound(
    tmp_path: Path,
) -> None:
    app = _app(tmp_path)
    try:
        client = app.test_client()
        tran_exchange = _pair(client, "tran")
        bearer = {"Authorization": f"Bearer {tran_exchange['token']}"}
        other_bearer = {
            "Authorization": f"Bearer {_pair(client, 'tran', 'outlook_addin')['token']}"
        }
        ngan_bearer = {"Authorization": f"Bearer {_pair(client, 'ngan')['token']}"}

        uploaded = client.post(
            "/api/companion/client/emails",
            data={
                "file": (
                    io.BytesIO(_lost_email()),
                    r"C:\private\Synthetic Companion.eml",
                    "message/rfc822",
                )
            },
            headers={**bearer, "X-Asset-Hub-Upload": "companion-email-v1"},
            content_type="multipart/form-data",
        )
        assert uploaded.status_code == 200
        upload_payload = uploaded.get_json()
        assert upload_payload["role"] == "tran"
        assert upload_payload["ingested"] == 1
        assert upload_payload["source_eml"]["filename"] == "Synthetic_Companion.eml"
        handle = upload_payload["source_eml"]["handle"]
        case_id = upload_payload["case_ids"][0]
        assert "download_url" not in str(upload_payload)

        dashboard_cases = client.get("/api/dashboard", headers=_BASIC).get_json()["cases"]
        source_case = next(case for case in dashboard_cases if case["id"] == case_id)
        row = source_case["metadata"]["asset_rows"][0]
        draft_request = {
            "assets": [
                {
                    "tag_number": row["asset_code"],
                    "asset_name": row["asset_name"],
                    "domain": row["domain"],
                    "lost_date": "2026-08-01",
                }
            ],
            "source_bindings": [{"case_id": case_id, "source_row_index": 0}],
            "mail_artifact_handle": handle,
            "body_intro": "Approved <script>alert(1)</script>",
            "processing_date": "2026-08-15",
        }
        created = client.post(
            "/api/tran/companion-drafts",
            json=draft_request,
            headers=_BASIC,
        )
        assert created.status_code == 201
        created_payload = created.get_json()
        package_id = created_payload["package_id"]
        assert created_payload["sent"] is False
        assert created_payload["package"]["source_eml_handle"] == handle

        unbound_list = client.get(
            "/api/companion/client/draft-packages", headers=other_bearer
        )
        assert unbound_list.status_code == 200
        assert unbound_list.get_json()["packages"] == []
        assert (
            client.get("/api/companion/client/draft-packages", headers=ngan_bearer).status_code
            == 403
        )

        listed = client.get("/api/companion/client/draft-packages", headers=bearer)
        assert listed.status_code == 200
        assert listed.headers["Cache-Control"] == "private, no-store"
        assert [item["id"] for item in listed.get_json()["packages"]] == [package_id]
        detail = client.get(
            f"/api/companion/client/draft-packages/{package_id}", headers=bearer
        )
        assert detail.status_code == 200
        package = detail.get_json()["package"]
        assert package["source_eml_handle"] == handle
        assert "<script>" not in package["body_html"]
        assert "&lt;script&gt;" in package["body_html"]
        workbook = base64.b64decode(package["workbook"]["content_base64"], validate=True)
        assert workbook.startswith(b"PK")

        bad_ack = client.post(
            f"/api/companion/client/draft-packages/{package_id}/ack",
            json={},
            headers=bearer,
        )
        assert bad_ack.status_code == 400
        acknowledged = client.post(
            f"/api/companion/client/draft-packages/{package_id}/ack",
            json={},
            headers={**bearer, "X-Asset-Hub-Action": "companion-ack-v1"},
        )
        assert acknowledged.get_json() == {
            "ok": True,
            "id": package_id,
            "acknowledged": True,
            "sent": False,
        }
        assert (
            client.get("/api/companion/client/draft-packages", headers=bearer)
            .get_json()["packages"]
            == []
        )

        cleared = client.post(
            "/api/test-data/clear",
            json={"confirm": "CLEAR_TEST_DATA"},
            headers={**_BASIC, "X-Asset-Hub-Action": "clear-test-data-v1"},
        )
        assert cleared.status_code == 200
        assert cleared.get_json()["companion_session_count"] == 3
        assert cleared.get_json()["companion_package_count"] == 1
        assert (
            client.get("/api/companion/client/draft-packages", headers=bearer).status_code
            == 401
        )
    finally:
        app.extensions["asset_hub"]["repository"].close()
