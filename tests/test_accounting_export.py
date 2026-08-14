from __future__ import annotations

import sys
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from asset_compensation.adapters import (
    AccountingValidationError,
    AccountingXlsxAdapter,
    OutputExistsError,
)


@dataclass
class SyntheticCase:
    id: str
    case_type: str
    domain: str
    asset_code: str
    amount: Decimal
    employee_name: str = "Synthetic User"
    asset_name: str = "Synthetic Device"
    received_at: datetime = datetime(2026, 8, 10, tzinfo=UTC)
    residual_value: Decimal | None = None
    responsibility_fee: Decimal | None = None
    repair_status: str | None = None
    supplier_number: str | None = "DEMO-SUPPLIER-01"
    supplier_site: str | None = "DEMO-SITE-A"
    warnings: tuple[str, ...] = ()
    source_file: str = "synthetic.eml"
    metadata: dict[str, Any] = field(default_factory=dict)


@pytest.fixture
def adapter() -> AccountingXlsxAdapter:
    return AccountingXlsxAdapter(
        {
            "DAMAGED": ("DEMO.DEBIT", "DEMO.DAMAGED.CREDIT"),
            "LOST": ("DEMO.DEBIT", "DEMO.LOST.CREDIT"),
        }
    )


def _cases() -> list[SyntheticCase]:
    return [
        SyntheticCase(
            id="case-damaged",
            case_type="DAMAGED",
            domain="=HYPERLINK(\"https://invalid.example\")",
            asset_code="DEMO-LAP-901",
            amount=Decimal("123456"),
            repair_status="NOT_REPAIRED",
        ),
        SyntheticCase(
            id="case-lost",
            case_type="LOST",
            domain="demo.mouse",
            asset_code="DEMO-MOU-902",
            amount=Decimal("246912"),
            residual_value=Decimal("234567"),
            responsibility_fee=Decimal("12345"),
        ),
    ]


def test_export_creates_balanced_auditable_workbook(
    tmp_path: Path, adapter: AccountingXlsxAdapter
) -> None:
    output = tmp_path / "accounting.xlsx"

    result = adapter.export(_cases(), output, expected_total=Decimal("370368"))

    assert result.path == output
    assert result.case_count == 2
    assert result.journal_line_count == 4
    assert result.source_total == Decimal("370368")
    assert result.debit_total == result.credit_total == result.source_total

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == ["Cases", "Journal", "Checks"]
        assert workbook["Cases"]["H4"].value == 123456
        assert workbook["Cases"]["H5"].value == 246912
        assert workbook["Cases"]["C4"].value.startswith("'=")
        assert workbook["Cases"]["C4"].data_type == "s"
        assert workbook["Journal"]["D4"].value == 123456
        assert workbook["Journal"]["E5"].value == 123456
        assert workbook["Checks"]["B4"].value == "=SUM('Journal'!D4:D7)"
        assert workbook["Checks"]["B5"].value == "=SUM('Journal'!E4:E7)"
        assert workbook["Checks"]["B7"].value == (
            "=MAX(ABS(B4-B5),ABS(B2-B4),ABS(B2-B5))"
        )
        assert workbook["Checks"]["B8"].value == '=IF(AND(B6=0,B7=0),"PASS","FAIL")'
        assert workbook["Cases"]["H4"].number_format == "#,##0"
    finally:
        workbook.close()


def test_export_does_not_overwrite_without_explicit_flag(
    tmp_path: Path, adapter: AccountingXlsxAdapter
) -> None:
    output = tmp_path / "accounting.xlsx"
    adapter.export(_cases(), output)

    with pytest.raises(OutputExistsError):
        adapter.export(_cases(), output)

    adapter.export(_cases(), output, overwrite=True)
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert workbook["Cases"]["A4"].value == "case-damaged"
        assert workbook["Checks"]["B8"].value == '=IF(AND(B6=0,B7=0),"PASS","FAIL")'
    finally:
        workbook.close()


def test_export_rejects_total_or_lost_breakdown_mismatch(
    tmp_path: Path, adapter: AccountingXlsxAdapter
) -> None:
    output = tmp_path / "invalid.xlsx"
    with pytest.raises(AccountingValidationError, match="expected total"):
        adapter.export(_cases(), output, expected_total=1)
    assert not output.exists()

    invalid = _cases()
    invalid[1].responsibility_fee = Decimal("1")
    with pytest.raises(AccountingValidationError, match="does not equal amount"):
        adapter.export(invalid, output)
    assert not output.exists()


def test_export_rejects_unsafe_gl_configuration() -> None:
    with pytest.raises(AccountingValidationError, match="cannot be the same"):
        AccountingXlsxAdapter({"DAMAGED": ("DEMO.GL", "DEMO.GL")})


def test_export_escapes_case_id_and_rejects_negative_components(
    tmp_path: Path, adapter: AccountingXlsxAdapter
) -> None:
    output = tmp_path / "safe.xlsx"
    malicious = _cases()
    malicious[0].id = "=1+1"

    adapter.export(malicious, output)

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["Cases"]["A4"].value == "'=1+1"
        assert workbook["Journal"]["A4"].value == "'=1+1"
    finally:
        workbook.close()

    invalid = _cases()
    invalid[1].residual_value = Decimal("-1")
    invalid[1].responsibility_fee = invalid[1].amount + 1
    with pytest.raises(AccountingValidationError, match="cannot be negative"):
        adapter.export(invalid, tmp_path / "negative.xlsx")


def test_pdf_module_does_not_import_pywin32_at_import_time() -> None:
    sys.modules.pop("win32com", None)
    sys.modules.pop("win32com.client", None)

    from asset_compensation.adapters import pdf

    assert pdf.WordPdfConverter is not None
    assert "win32com.client" not in sys.modules


def test_pdf_atomic_commit_never_clobbers_a_racing_writer(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from asset_compensation.adapters import pdf

    temporary = tmp_path / "unique-temporary.pdf"
    destination = tmp_path / "destination.pdf"
    temporary.write_bytes(b"synthetic-new-pdf")

    def create_competing_destination(source: object, target: object) -> None:
        del source
        Path(target).write_bytes(b"synthetic-race-winner")
        raise FileExistsError

    monkeypatch.setattr(pdf.os, "link", create_competing_destination)
    with pytest.raises(FileExistsError, match="already exists"):
        pdf._commit_temp_file(temporary, destination, overwrite=False)

    assert destination.read_bytes() == b"synthetic-race-winner"
    assert temporary.read_bytes() == b"synthetic-new-pdf"
