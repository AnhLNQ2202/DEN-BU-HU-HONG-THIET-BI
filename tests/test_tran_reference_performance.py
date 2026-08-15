"""Structural and cache regressions for TranNNB reference performance."""

from __future__ import annotations

import io
import os
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path
from threading import Event, Lock

import pytest
from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from asset_compensation.adapters import (
    CcdcWorkbookIndex,
    FaGlWorkbookIndex,
)
from asset_compensation.domain import ReferenceStatus
from asset_compensation.services.tran_reference_upload_service import (
    TranReferenceUpload,
    TranReferenceUploadError,
    TranReferenceUploadService,
)

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FA_SHEETS = {
    "VNG-Asset": (3, 4),
    "VNG-Tool": (3, 4),
    "VNGS-Asset": (2, 3),
    "VNGS-Tool": (2, 3),
}
_FA_HEADERS = {
    2: "Company Name",
    7: "Cost Center",
    8: "Product Code",
    10: "Location",
    12: "Asset",
    16: "Tag Number",
    22: "Date of Depreciation",
    25: "Life(month)",
    27: "Cost",
}


def _workbook_bytes(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _fa_gl_bytes(tag_number: str = "MOU10001") -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, (header_row, data_row) in _FA_SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        for column, header in _FA_HEADERS.items():
            sheet.cell(header_row, column, header)
        if sheet_name == "VNG-Tool":
            sheet.cell(data_row, 2, "Synthetic Company")
            sheet.cell(data_row, 7, "0603")
            sheet.cell(data_row, 8, "000")
            sheet.cell(data_row, 10, "01")
            sheet.cell(data_row, 12, f"ASSET-{tag_number}")
            sheet.cell(data_row, 16, tag_number)
            sheet.cell(data_row, 22, date(2025, 1, 1))
            sheet.cell(data_row, 25, 48)
            sheet.cell(data_row, 27, 1_000_000)
    return _workbook_bytes(workbook)


def _ccdc_bytes() -> bytes:
    workbook = Workbook()
    define = workbook.active
    define.title = "Define"
    define.append(["Product Type", "Barcode", "Group Type"])
    define.append(["Other/Spe.part", "MOU", "Hardware"])
    warehouse = workbook.create_sheet("BC Xuatkho")
    warehouse.append(["Unused", "Asset Name", "Unused", "Unused", "Start time"])
    warehouse.append([None, "MOU10001", None, None, date(2025, 1, 1)])
    return _workbook_bytes(workbook)


def _upload(payload: bytes, filename: str) -> TranReferenceUpload:
    return TranReferenceUpload(
        filename=filename,
        content_type=_XLSX_MIME,
        stream=io.BytesIO(payload),
    )


def _write_with_new_signature(path: Path, payload: bytes) -> None:
    previous_mtime = path.stat().st_mtime_ns if path.exists() else 0
    path.write_bytes(payload)
    timestamp = max(time.time_ns(), previous_mtime + 1_000_000)
    os.utime(path, ns=(timestamp, timestamp))


def test_ccdc_read_only_scan_never_uses_random_cell_access(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Read-only sheets must be streamed once instead of reparsed per cell."""

    path = tmp_path / "sparse-ccdc.xlsx"
    workbook = Workbook()
    define = workbook.active
    define.title = "Define"
    define.cell(3, 1, "Product Type")
    define.cell(3, 2, "Barcode")
    define.cell(3, 3, "Group Type")
    define.cell(5, 1, "Computer asset")
    define.cell(5, 2, "LAP")
    define.cell(5, 3, "Hardware")

    cmdb = workbook.create_sheet("CMDB")
    cmdb.cell(4, 1, "Asset Name")
    cmdb.cell(4, 2, "Product Type")
    cmdb.cell(6, 1, "OLD10001")
    cmdb.cell(6, 2, "Computer component asset")

    warehouse = workbook.create_sheet("BC Xuatkho")
    warehouse.cell(1, 1, "Legacy export without recognized headers")
    warehouse.cell(3, 2, "ADA10001")
    warehouse.cell(3, 5, date(2024, 5, 1))
    warehouse.cell(7, 2, "ADA10001")
    warehouse.cell(7, 5, date(2023, 7, 2))
    workbook.save(path)
    workbook.close()

    def fail_random_access(*_args: object, **_kwargs: object) -> object:
        raise AssertionError("read-only CCDC parsing must use iter_rows")

    monkeypatch.setattr(ReadOnlyWorksheet, "cell", fail_random_access)
    index = CcdcWorkbookIndex.from_path(path)

    laptop = index.classification("LAP")
    fallback = index.classification("OLD")
    assert laptop.status is ReferenceStatus.MATCHED
    assert laptop.classification is not None
    assert laptop.classification.source_sheet == "Define"
    assert laptop.classification.source_row == 5
    assert fallback.status is ReferenceStatus.MATCHED
    assert fallback.classification is not None
    assert fallback.classification.source_sheet == "CMDB"
    assert fallback.classification.source_row == 6
    assert index.earliest_start_date("ADA10001") == date(2023, 7, 2)


def test_successful_upload_primes_exact_validated_indexes_and_replaces_generation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    real_fa_loader = FaGlWorkbookIndex.from_path.__func__
    real_ccdc_loader = CcdcWorkbookIndex.from_path.__func__
    fa_indexes: list[FaGlWorkbookIndex] = []
    ccdc_indexes: list[CcdcWorkbookIndex] = []

    def track_fa(cls: type[FaGlWorkbookIndex], path: str | Path) -> FaGlWorkbookIndex:
        index = real_fa_loader(cls, path)
        fa_indexes.append(index)
        return index

    def track_ccdc(cls: type[CcdcWorkbookIndex], path: str | Path) -> CcdcWorkbookIndex:
        index = real_ccdc_loader(cls, path)
        ccdc_indexes.append(index)
        return index

    monkeypatch.setattr(FaGlWorkbookIndex, "from_path", classmethod(track_fa))
    monkeypatch.setattr(CcdcWorkbookIndex, "from_path", classmethod(track_ccdc))

    service.upload(
        _upload(_fa_gl_bytes("MOU10001"), "fa.xlsx"),
        _upload(_ccdc_bytes(), "ccdc.xlsx"),
    )
    first_fa, first_ccdc = service.load_indices()
    assert first_fa is fa_indexes[0]
    assert first_ccdc is ccdc_indexes[0]
    assert len(fa_indexes) == len(ccdc_indexes) == 1
    assert first_fa.source_path.parent.parent == service.versions_dir
    assert first_ccdc is not None
    assert first_ccdc.source_path.parent.parent == service.versions_dir

    service.upload(
        _upload(_fa_gl_bytes("MOU20002"), "fa.xlsx"),
        clear_ccdc=True,
    )
    second_fa, second_ccdc = service.load_indices()
    assert second_fa is fa_indexes[1]
    assert second_fa is not first_fa
    assert second_fa.lookup("MOU20002").status is ReferenceStatus.MATCHED
    assert second_ccdc is None
    assert len(fa_indexes) == 2
    assert len(ccdc_indexes) == 1


def test_failed_upload_keeps_prior_index_cache(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    service.upload(_upload(_fa_gl_bytes(), "fa.xlsx"))
    before = service.load_indices()

    with pytest.raises(TranReferenceUploadError):
        service.upload(_upload(b"not-an-xlsx", "fa.xlsx"))

    after = service.load_indices()
    assert after[0] is before[0]
    assert after[1] is before[1]


def test_clear_invalidates_managed_cache_and_allows_external_fallback(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    service.upload(_upload(_fa_gl_bytes(), "fa.xlsx"))
    service.load_indices()
    service.clear()

    with pytest.raises(TranReferenceUploadError, match="not configured"):
        service.load_indices()

    external = tmp_path / "external-fa.xlsx"
    external.write_bytes(_fa_gl_bytes("MOU30003"))
    fa_gl, ccdc = service.load_indices(external, None)
    assert fa_gl.lookup("MOU30003").status is ReferenceStatus.MATCHED
    assert ccdc is None


def test_external_cache_reloads_only_after_stable_stat_signature_change(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    external = tmp_path / "external-fa.xlsx"
    external.write_bytes(_fa_gl_bytes("MOU10001"))
    real_loader = FaGlWorkbookIndex.from_path.__func__
    calls = 0

    def tracked(cls: type[FaGlWorkbookIndex], path: str | Path) -> FaGlWorkbookIndex:
        nonlocal calls
        calls += 1
        return real_loader(cls, path)

    monkeypatch.setattr(FaGlWorkbookIndex, "from_path", classmethod(tracked))
    first, _ = service.load_indices(external, None)
    repeated, _ = service.load_indices(external, None)
    assert repeated is first
    assert calls == 1

    _write_with_new_signature(external, _fa_gl_bytes("MOU20002"))
    changed, _ = service.load_indices(external, None)
    stable, _ = service.load_indices(external, None)
    assert changed is stable
    assert changed is not first
    assert changed.lookup("MOU20002").status is ReferenceStatus.MATCHED
    assert calls == 2


def test_external_mutation_during_index_build_is_rejected(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    external = tmp_path / "external-fa.xlsx"
    external.write_bytes(_fa_gl_bytes("MOU10001"))
    replacement = _fa_gl_bytes("MOU20002")
    real_loader = FaGlWorkbookIndex.from_path.__func__

    def mutate_after_load(
        cls: type[FaGlWorkbookIndex],
        path: str | Path,
    ) -> FaGlWorkbookIndex:
        index = real_loader(cls, path)
        _write_with_new_signature(Path(path), replacement)
        return index

    monkeypatch.setattr(
        FaGlWorkbookIndex,
        "from_path",
        classmethod(mutate_after_load),
    )
    with pytest.raises(TranReferenceUploadError, match="changed while"):
        service.load_indices(external, None)

    monkeypatch.setattr(
        FaGlWorkbookIndex,
        "from_path",
        classmethod(real_loader),
    )
    recovered, _ = service.load_indices(external, None)
    assert recovered.lookup("MOU20002").status is ReferenceStatus.MATCHED


def test_external_pair_replacement_between_loads_never_publishes_mixed_indexes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    fa_gl = tmp_path / "fa.xlsx"
    ccdc = tmp_path / "ccdc.xlsx"
    fa_gl.write_bytes(_fa_gl_bytes("MOU10001"))
    ccdc.write_bytes(_ccdc_bytes())

    old_fa, old_ccdc = service.load_indices(fa_gl, ccdc)
    assert old_fa.lookup("MOU10001").status is ReferenceStatus.MATCHED
    assert old_ccdc is not None

    real_load = service._load_cached_index
    calls = 0

    def replace_pair_after_first_load(**kwargs: object) -> tuple[object, object]:
        nonlocal calls
        result = real_load(**kwargs)
        calls += 1
        if calls == 1:
            _write_with_new_signature(fa_gl, _fa_gl_bytes("MOU20002"))
            _write_with_new_signature(ccdc, _ccdc_bytes())
        return result

    monkeypatch.setattr(service, "_load_cached_index", replace_pair_after_first_load)
    with pytest.raises(TranReferenceUploadError, match="changed while"):
        service.load_indices(fa_gl, ccdc)

    # The failed pair is not published; a stable retry produces one coherent
    # replacement generation instead of reusing either half of the old pair.
    monkeypatch.setattr(service, "_load_cached_index", real_load)
    new_fa, new_ccdc = service.load_indices(fa_gl, ccdc)
    assert new_fa is not old_fa
    assert new_fa.lookup("MOU20002").status is ReferenceStatus.MATCHED
    assert new_ccdc is not old_ccdc


def test_concurrent_external_cache_miss_builds_one_index(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = TranReferenceUploadService(tmp_path / "reference")
    external = tmp_path / "external-fa.xlsx"
    external.write_bytes(_fa_gl_bytes())
    real_loader = FaGlWorkbookIndex.from_path.__func__
    started = Event()
    release = Event()
    count_lock = Lock()
    calls = 0

    def paused(cls: type[FaGlWorkbookIndex], path: str | Path) -> FaGlWorkbookIndex:
        nonlocal calls
        with count_lock:
            calls += 1
        started.set()
        assert release.wait(timeout=5)
        return real_loader(cls, path)

    monkeypatch.setattr(FaGlWorkbookIndex, "from_path", classmethod(paused))
    with ThreadPoolExecutor(max_workers=2) as executor:
        first_future = executor.submit(service.load_indices, external, None)
        assert started.wait(timeout=5)
        second_future = executor.submit(service.load_indices, external, None)
        release.set()
        first = first_future.result(timeout=5)[0]
        second = second_future.result(timeout=5)[0]

    assert first is second
    assert calls == 1
