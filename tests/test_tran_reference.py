"""Synthetic reference-workbook tests for the TranNNB workflow."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from asset_compensation.adapters import (
    CcdcWorkbookIndex,
    FaGlWorkbookIndex,
    TranReferenceError,
)
from asset_compensation.domain import DepreciationGroup, ReferenceStatus

FA_SHEETS = {
    "VNG-Asset": (3, 4),
    "VNG-Tool": (3, 4),
    "VNGS-Asset": (2, 3),
    "VNGS-Tool": (2, 3),
}


def make_fa_gl(path: Path, rows: dict[str, list[dict[str, object]]]) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, (header_row, data_row) in FA_SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        headers = {
            2: "Company Name",
            7: "Cost Center",
            8: "Product Code",
            10: "Location",
            12: "Asset",
            16: "Tag Number",
            22: "Date of Depreciation",
            25: "Life(month)",
            27: "Cost",
        }
        for column, header in headers.items():
            sheet.cell(header_row, column, header)
        for offset, record in enumerate(rows.get(sheet_name, [])):
            row = data_row + offset
            sheet.cell(row, 2, record.get("company", "Synthetic Company"))
            sheet.cell(row, 7, record.get("cost_center", "0603"))
            sheet.cell(row, 8, record.get("product_code", "000"))
            sheet.cell(row, 10, record.get("location", "01"))
            sheet.cell(row, 12, record.get("asset_number", "SYN-001"))
            sheet.cell(row, 16, record["tag"])
            sheet.cell(row, 22, record.get("start_date", date(2025, 1, 1)))
            sheet.cell(row, 25, 36)
            sheet.cell(row, 27, record.get("cost", 1_000_000))
    workbook.save(path)
    workbook.close()
    return path


def make_ccdc(path: Path) -> Path:
    workbook = Workbook()
    define = workbook.active
    define.title = "Define"
    define.append(["Product Type", "Barcode", "Group Type"])
    define.append(["Computer asset", "LAP", "Hardware"])
    define.append(["Other/Spe.part", "ZZZ", "Hardware"])
    define.append(["Cloud product", "SRV", "Service"])
    cmdb = workbook.create_sheet("CMDB")
    cmdb.append(["Asset Name", "Product Type"])
    cmdb.append(["OLD10001", "Computer component asset"])
    warehouse = workbook.create_sheet("BC Xuatkho")
    warehouse.append(["Unused", "Asset Name", "Unused", "Unused", "Start time"])
    warehouse.append([None, "ADA10001", None, None, date(2024, 5, 1)])
    warehouse.append([None, "ADA10001", None, None, date(2023, 7, 2)])
    workbook.save(path)
    workbook.close()
    return path


def test_fa_gl_uses_exact_sheet_offsets_and_reports_ambiguous_matches(tmp_path: Path) -> None:
    path = make_fa_gl(
        tmp_path / "synthetic-fa-gl.xlsx",
        {
            "VNG-Asset": [
                {
                    "tag": "LAP10001",
                    "asset_number": "A-001",
                    "start_date": date(2024, 2, 3),
                    "cost": 2_000_000,
                },
                {"tag": "DUP10001"},
            ],
            "VNGS-Tool": [{"tag": "DUP10001", "asset_number": "T-002"}],
        },
    )

    index = FaGlWorkbookIndex.from_path(path)
    found = index.lookup(" lap10001 ")
    duplicate = index.lookup("DUP10001")
    missing = index.lookup("NONE")

    assert found.status is ReferenceStatus.MATCHED
    assert found.record is not None
    assert found.record.asset_number == "A-001"
    assert found.record.book == "Asset"
    assert found.record.entity == "VNG"
    assert found.record.start_date == date(2024, 2, 3)
    assert found.record.cost == 2_000_000
    assert "cột P" in found.record.source_note
    assert duplicate.status is ReferenceStatus.AMBIGUOUS
    assert len(duplicate.matches) == 2
    assert missing.status is ReferenceStatus.NOT_FOUND


def test_fa_gl_rejects_missing_required_sheets(tmp_path: Path) -> None:
    workbook = Workbook()
    path = tmp_path / "bad-fa-gl.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(TranReferenceError, match="missing required sheets"):
        FaGlWorkbookIndex.from_path(path)


def test_ccdc_reads_define_cmdb_and_oldest_warehouse_date(tmp_path: Path) -> None:
    index = CcdcWorkbookIndex.from_path(make_ccdc(tmp_path / "synthetic-ccdc.xlsx"))

    laptop = index.classification("LAP")
    fallback = index.classification("OLD")
    service = index.classification("SRV")

    assert laptop.status is ReferenceStatus.MATCHED
    assert laptop.classification is not None
    assert laptop.classification.group is DepreciationGroup.SIX_YEAR
    assert laptop.classification.physical is True
    assert fallback.classification is not None
    assert fallback.classification.source_sheet == "CMDB"
    assert fallback.classification.group is DepreciationGroup.FOUR_YEAR
    assert service.classification is not None
    assert service.classification.physical is False
    assert index.earliest_start_date("ADA10001") == date(2023, 7, 2)

