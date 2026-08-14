"""HTTP contract tests for the local dashboard."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from asset_compensation.adapters import ACCOUNTING_TEMPLATE_HEADERS
from asset_compensation.config import Settings
from asset_compensation.domain import CaseType, ParsedCase
from asset_compensation.web import create_app


@pytest.fixture
def app(tmp_path: Path) -> Flask:
    application = create_app(
        Settings(
            data_dir=tmp_path / "runtime",
            demo_mode=True,
            secret_key="synthetic-test-key",
        )
    )
    application.config.update(TESTING=True)
    yield application
    application.extensions["asset_hub"]["repository"].close()


def test_dashboard_serves_synthetic_cases_and_security_headers(app: Flask) -> None:
    client = app.test_client()

    page = client.get("/")
    dashboard = client.get("/api/dashboard")

    assert page.status_code == 200
    assert b"Asset Compensation Hub" in page.data
    assert "unsafe-inline" not in page.headers["Content-Security-Policy"]
    assert dashboard.status_code == 200
    payload = dashboard.get_json()
    assert payload["ok"] is True
    assert payload["summary"]["total"] == 5
    assert all(case["domain"].startswith("demo.") for case in payload["cases"])


@pytest.mark.parametrize(
    "query",
    ["limit=abc", "limit=0", "offset=-1", "status=BOGUS", "has_warnings=perhaps"],
)
def test_case_filters_reject_invalid_values_as_json(app: Flask, query: str) -> None:
    response = app.test_client().get(f"/api/cases?{query}")

    assert response.status_code == 400
    assert response.is_json
    assert response.get_json()["ok"] is False


def test_batch_export_is_balanced_downloadable_and_retry_safe(app: Flask) -> None:
    client = app.test_client()
    initial = client.get("/api/dashboard").get_json()
    ready_ids = [
        case["id"]
        for case in initial["cases"]
        if case["status"] == "READY_FOR_ACCOUNTING"
    ][:2]
    request_body = {"batch_name": "GN2010126", "case_ids": ready_ids}

    created = client.post("/api/batches", json=request_body)

    assert created.status_code == 201
    batch = created.get_json()["batch"]
    assert batch["case_count"] == 2
    assert batch["download_url"]

    downloaded = client.get(batch["download_url"])
    assert downloaded.status_code == 200
    assert downloaded.headers["Cache-Control"] == "private, no-store"
    assert downloaded.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    output_name = batch["metadata"]["output_name"]
    output_path = app.extensions["asset_hub"]["settings"].output_dir / output_name
    workbook = load_workbook(output_path, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Sheet1"]
        sheet = workbook["Sheet1"]
        assert tuple(sheet.cell(1, column).value for column in range(1, 31)) == (
            ACCOUNTING_TEMPLATE_HEADERS
        )
        assert [sheet.cell(row, 1).value for row in range(2, 6)] == [
            "Prepayment",
            "Credit Memo",
            "Prepayment",
            "Credit Memo",
        ]
        assert sum(sheet.cell(row, 7).value for row in range(2, 6)) == 0
        assert batch["metadata"]["format"] == "accounting-template-v1"
        assert output_name == "hachtoan_gop_01.2026_GN2010126.xlsx"
    finally:
        workbook.close()

    retried = client.post("/api/batches", json=request_body)
    assert retried.status_code == 200
    assert retried.get_json()["batch"]["id"] == batch["id"]


def test_batch_name_must_contain_a_real_date(app: Flask) -> None:
    response = app.test_client().post(
        "/api/batches",
        json={"batch_name": "GN2310226", "case_ids": ["synthetic-id"]},
    )

    assert response.status_code == 400
    assert "invalid calendar date" in response.get_json()["error"]


def test_batch_uses_bounded_invoice_start_and_semantic_policy_resolution(
    app: Flask,
) -> None:
    client = app.test_client()
    ready_damaged = next(
        case
        for case in client.get("/api/dashboard").get_json()["cases"]
        if case["status"] == "READY_FOR_ACCOUNTING" and case["case_type"] == "DAMAGED"
    )

    created = client.post(
        "/api/batches",
        json={
            "batch_name": "GN2020226",
            "case_ids": [ready_damaged["id"]],
            "invoice_start": 0,
        },
    )

    assert created.status_code == 201
    batch = created.get_json()["batch"]
    assert batch["metadata"]["invoice_start"] == 0
    assert "ASSET_COMPENSATION_PREPAYMENT" in batch["metadata"][
        "accounting_policy_keys"
    ]
    downloaded = client.get(batch["download_url"])
    workbook = load_workbook(io.BytesIO(downloaded.data), read_only=True)
    try:
        assert workbook.active["B2"].value == "GN2020226000"
        assert workbook.active["B3"].value == "GN2020226001"
    finally:
        workbook.close()


def test_batch_rejects_unverified_repair_policy_before_writing_output(app: Flask) -> None:
    extension = app.extensions["asset_hub"]
    service = extension["case_service"]
    candidate = next(
        case for case in service.list_cases() if case.asset_code == "DEMO-PHO-009"
    )
    service.transition_status(candidate.id, "READY_FOR_ACCOUNTING", actor="synthetic-test")
    before = tuple(extension["settings"].output_dir.iterdir())

    response = app.test_client().post(
        "/api/batches",
        json={"batch_name": "GN2030226", "case_ids": [candidate.id]},
    )

    assert response.status_code == 400
    assert "repair status" in response.get_json()["error"]
    assert tuple(extension["settings"].output_dir.iterdir()) == before


def test_batch_rejects_zero_value_case_without_publishing_output(app: Flask) -> None:
    extension = app.extensions["asset_hub"]
    service = extension["case_service"]
    candidate = service.ingest_one(
        ParsedCase(
            case_type=CaseType.LOST,
            domain="demo.zero",
            employee_name="Synthetic Zero",
            asset_code="DEMO-ZERO-001",
            asset_name="Synthetic zero-value asset",
            received_at=datetime(2026, 8, 14, tzinfo=UTC),
            amount=0,
            residual_value=0,
            responsibility_fee=0,
            supplier_number="SYN-000",
            supplier_site="DEMO",
            supplier_name="Synthetic Supplier",
            source_file="synthetic-zero.eml",
            source_id="synthetic-zero@example.invalid",
            metadata={},
        )
    )
    service.transition_status(
        candidate.id,
        "READY_FOR_ACCOUNTING",
        actor="synthetic-test",
    )
    before = tuple(extension["settings"].output_dir.iterdir())

    response = app.test_client().post(
        "/api/batches",
        json={"batch_name": "GN2040226", "case_ids": [candidate.id]},
    )

    assert response.status_code == 400
    assert "greater than zero" in response.get_json()["error"]
    assert tuple(extension["settings"].output_dir.iterdir()) == before


def test_optional_shared_demo_auth_keeps_health_check_public(tmp_path: Path) -> None:
    protected = create_app(
        Settings(
            data_dir=tmp_path / "protected-runtime",
            access_user="demo-team",
            access_password="synthetic-secret",
        )
    )
    protected.config.update(TESTING=True)
    client = protected.test_client()
    try:
        health = client.get("/api/health")
        assert health.status_code == 200
        assert health.get_json() == {
            "ok": True,
            "service": "asset-compensation-hub",
        }
        denied = client.get("/")
        assert denied.status_code == 401
        assert denied.headers["WWW-Authenticate"].startswith("Basic")
        assert (
            client.get(
                "/",
                headers={"Authorization": "Basic ZGVtby10ZWFtOnN5bnRoZXRpYy1zZWNyZXQ="},
            ).status_code
            == 200
        )
    finally:
        protected.extensions["asset_hub"]["repository"].close()
