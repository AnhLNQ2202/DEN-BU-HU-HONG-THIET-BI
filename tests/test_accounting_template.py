from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from asset_compensation.adapters import (
    ACCOUNTING_TEMPLATE_HEADERS,
    AccountingTemplateAdapter,
    AccountingValidationError,
    OutputExistsError,
)


@dataclass
class TemplateCase:
    id: str
    case_type: str
    domain: str
    asset_code: str
    amount: object
    supplier_number: object = "000101"
    supplier_site: object = "2201"
    supplier_name: object = "Synthetic Supplier"
    residual_value: object | None = None
    responsibility_fee: object | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@pytest.fixture
def adapter() -> AccountingTemplateAdapter:
    return AccountingTemplateAdapter(
        {
            "DAMAGED": ("DEMO.DEBIT", "DEMO.DAMAGED.CREDIT"),
            "LOST": ("DEMO.DEBIT", "DEMO.LOST.CREDIT"),
        },
        org_id="99",
    )


def _cases() -> list[TemplateCase]:
    return [
        TemplateCase(
            id="demo-case-damaged",
            case_type="DAMAGED",
            domain="demo.user",
            asset_code="DEMO-LAP-901",
            amount=Decimal("123456"),
        ),
        TemplateCase(
            id="demo-case-lost",
            case_type="LOST",
            domain="demo.mouse",
            asset_code="DEMO-MOU-902",
            amount=Decimal("246912"),
            supplier_number="990102",
            supplier_site="0022",
            residual_value=Decimal("234567"),
            responsibility_fee=Decimal("12345"),
            metadata={
                "credit_lines": [
                    {"account": "DEMO.LOST.VALUE", "amount": Decimal("234567")},
                    {"account": "DEMO.LOST.FEE", "amount": Decimal("12345")},
                ]
            },
        ),
    ]


def test_export_maps_invoices_and_preserves_template_styles(
    tmp_path: Path,
    adapter: AccountingTemplateAdapter,
) -> None:
    output_stem = tmp_path / "synthetic-accounting"
    export_date = date(2026, 8, 14)
    template_bytes = adapter.template_path.read_bytes()

    result = adapter.export(
        _cases(),
        output_stem,
        batch_name="DEMO-BATCH-",
        invoice_date=export_date,
        expected_total=Decimal("370368"),
    )

    assert result.path == output_stem.with_suffix(".xlsx")
    assert result.template_path == adapter.template_path
    assert result.output_suffix == adapter.output_suffix == ".xlsx"
    assert result.case_count == 2
    assert result.invoice_count == 4
    assert result.row_count == result.journal_line_count == 5
    assert result.source_total == result.credit_total == Decimal("370368")
    assert result.invoice_numbers == (
        "DEMO-BATCH-001",
        "DEMO-BATCH-002",
        "DEMO-BATCH-003",
        "DEMO-BATCH-004",
    )
    assert adapter.template_path.read_bytes() == template_bytes

    template = load_workbook(adapter.template_path)
    workbook = load_workbook(result.path, data_only=False)
    try:
        source_sheet = template.active
        sheet = workbook.active
        assert tuple(sheet.cell(1, column).value for column in range(1, 31)) == (
            ACCOUNTING_TEMPLATE_HEADERS
        )
        assert sheet["A2"].value == "Prepayment"
        assert sheet["B2"].value == "DEMO-BATCH-001"
        assert sheet["C2"].value == "000101"
        assert sheet["C2"].data_type == "s"
        assert sheet["D2"].value == 2201
        assert sheet["E2"].value.date() == export_date
        assert sheet["F2"].value.date() == export_date
        assert sheet["R2"].value is None
        assert sheet["AD2"].value is None
        assert sheet["G2"].value == 123456
        assert sheet["L2"].value == 123456
        assert sheet["M2"].value == "DEMO.DEBIT"
        assert sheet["W2"].value == "DEMO-BATCH-"
        assert sheet["Y2"].value == 99
        assert sheet["I2"].value == (
            "Trừ lương demo.user đền bù do hư hỏng tài sản "
            "DEMO-LAP-901, chưa xác định tình trạng sửa chữa"
        )

        assert sheet["A3"].value == "Credit Memo"
        assert sheet["B3"].value == "DEMO-BATCH-002"
        assert sheet["G3"].value == -123456
        assert sheet["L3"].value == -123456
        assert sheet["M3"].value == "DEMO.DAMAGED.CREDIT"

        assert sheet["A4"].value == "Prepayment"
        assert sheet["B4"].value == "DEMO-BATCH-003"
        assert sheet["C4"].value == 990102
        assert sheet["D4"].value == "0022"
        assert sheet["D4"].data_type == "s"
        assert sheet["G4"].value == 246912
        assert sheet["M4"].value == "DEMO.DEBIT"
        assert sheet["I4"].value == (
            "Trừ lương demo.mouse đền bù do thất lạc tài sản DEMO-MOU-902"
        )

        assert sheet["A5"].value == sheet["A6"].value == "Credit Memo"
        assert sheet["B5"].value == sheet["B6"].value == "DEMO-BATCH-004"
        assert sheet["G5"].value == sheet["G6"].value == -246912
        assert sheet["K5"].value == 1
        assert sheet["K6"].value == 2
        assert sheet["L5"].value == -234567
        assert sheet["L6"].value == -12345
        assert sheet["M5"].value == "DEMO.LOST.VALUE"
        assert sheet["M6"].value == "DEMO.LOST.FEE"
        assert all(sheet.cell(7, column).value is None for column in range(1, 31))

        for column in range(1, 31):
            assert sheet.cell(2, column).style_id == source_sheet.cell(2, column).style_id
            assert sheet.cell(3, column).style_id == source_sheet.cell(3, column).style_id
            assert sheet.cell(4, column).style_id == source_sheet.cell(2, column).style_id
            assert sheet.cell(5, column).style_id == source_sheet.cell(3, column).style_id
            assert sheet.cell(6, column).style_id == source_sheet.cell(3, column).style_id
        assert sheet.column_dimensions["A"].width == source_sheet.column_dimensions["A"].width
        assert len(sheet.conditional_formatting) == len(source_sheet.conditional_formatting)
    finally:
        workbook.close()
        template.close()


def test_custom_credit_type_tables_and_styles_survive_but_sample_values_do_not(
    tmp_path: Path,
) -> None:
    built_in = AccountingTemplateAdapter(
        {"LOST": ("DEMO.DEBIT", "DEMO.CREDIT")}
    ).template_path
    custom_template = tmp_path / "synthetic-template.xlsx"
    shutil.copyfile(built_in, custom_template)
    workbook = load_workbook(custom_template)
    try:
        sheet = workbook.active
        sheet["A3"] = "Synthetic Reversal"
        sheet["H2"] = "REMOVE-PREPAY-DEFAULT"
        sheet["H3"] = "REMOVE-CREDIT-DEFAULT"
        for row in range(4, 8):
            sheet.cell(row, 8, f"REMOVE-SAMPLE-{row}")
        table = Table(displayName="SyntheticImport", ref="A1:AD7")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        workbook.save(custom_template)
    finally:
        workbook.close()

    case = _cases()[1]
    adapter = AccountingTemplateAdapter(
        {"LOST": ("DEMO.DEBIT", "DEMO.CREDIT")},
        template_path=custom_template,
    )
    result = adapter.export(
        [case],
        tmp_path / "custom-output.xlsx",
        batch_name="DEMO-CUSTOM-",
        invoice_date=date(2026, 8, 14),
    )

    output = load_workbook(result.path)
    try:
        sheet = output.active
        assert sheet["A3"].value == sheet["A4"].value == "Synthetic Reversal"
        assert sheet["H2"].value is None
        assert sheet["H3"].value is None
        assert sheet["H4"].value is None
        assert all(sheet.cell(row, 8).value is None for row in range(5, 8))
        assert sheet.tables["SyntheticImport"].ref == "A1:AD4"
    finally:
        output.close()


def test_formula_injection_is_escaped_in_every_mapped_text_field(tmp_path: Path) -> None:
    case = TemplateCase(
        id="demo-formula-case",
        case_type="DAMAGED",
        domain="demo.formula",
        asset_code="DEMO-SAFE-904",
        amount=Decimal("123456"),
        supplier_number="+DANGEROUS",
        supplier_site="-DANGEROUS",
        supplier_name="@DANGEROUS",
        metadata={"description": "=DANGEROUS"},
    )
    adapter = AccountingTemplateAdapter(
        {"DAMAGED": ("DEMO.DEBIT", "DEMO.CREDIT")}
    )

    result = adapter.export(
        [case],
        tmp_path / "safe.xlsx",
        batch_name="=DANGEROUS",
        invoice_date=date(2026, 8, 14),
    )

    workbook = load_workbook(result.path, data_only=False)
    try:
        sheet = workbook.active
        for coordinate in ("B2", "C2", "D2", "I2", "J2", "P2", "W2"):
            assert str(sheet[coordinate].value).startswith("'")
            assert sheet[coordinate].data_type == "s"
    finally:
        workbook.close()


@pytest.mark.parametrize("amount", [Decimal("-1"), Decimal("1.5"), True])
def test_rejects_non_whole_or_negative_vnd(tmp_path: Path, amount: object) -> None:
    case = TemplateCase(
        id="demo-invalid-amount",
        case_type="DAMAGED",
        domain="demo.invalid",
        asset_code="DEMO-INVALID-905",
        amount=amount,
    )
    adapter = AccountingTemplateAdapter(
        {"DAMAGED": ("DEMO.DEBIT", "DEMO.CREDIT")}
    )

    with pytest.raises(AccountingValidationError, match="whole VND|negative"):
        adapter.export(
            [case],
            tmp_path / "invalid.xlsx",
            batch_name="DEMO-INVALID-",
            invoice_date=date(2026, 8, 14),
        )


def test_rejects_invalid_headers_credit_totals_and_expected_total(tmp_path: Path) -> None:
    base_adapter = AccountingTemplateAdapter(
        {
            "DAMAGED": ("DEMO.DEBIT", "DEMO.CREDIT"),
            "LOST": ("DEMO.DEBIT", "DEMO.LOST.CREDIT"),
        }
    )
    bad_template = tmp_path / "bad-template.xlsx"
    shutil.copyfile(base_adapter.template_path, bad_template)
    workbook = load_workbook(bad_template)
    try:
        workbook.active["A1"] = "Unexpected Header"
        workbook.save(bad_template)
    finally:
        workbook.close()
    bad_adapter = AccountingTemplateAdapter(
        {"DAMAGED": ("DEMO.DEBIT", "DEMO.CREDIT")},
        template_path=bad_template,
    )
    with pytest.raises(AccountingValidationError, match="30-column"):
        bad_adapter.export(
            [_cases()[0]],
            tmp_path / "bad-header.xlsx",
            batch_name="DEMO-BAD-",
            invoice_date=date(2026, 8, 14),
        )

    mismatched = _cases()[0]
    mismatched.metadata["credit_lines"] = [("DEMO.CREDIT", Decimal("1"))]
    with pytest.raises(AccountingValidationError, match="do not reconcile"):
        base_adapter.export(
            [mismatched],
            tmp_path / "bad-credit.xlsx",
            batch_name="DEMO-BAD-",
            invoice_date=date(2026, 8, 14),
        )
    with pytest.raises(AccountingValidationError, match="expected total"):
        base_adapter.export(
            [_cases()[0]],
            tmp_path / "bad-total.xlsx",
            batch_name="DEMO-BAD-",
            invoice_date=date(2026, 8, 14),
            expected_total=Decimal("123457"),
        )

    lost_mismatch = _cases()[1]
    lost_mismatch.responsibility_fee = Decimal("1")
    with pytest.raises(AccountingValidationError, match="residual value"):
        base_adapter.export(
            [lost_mismatch],
            tmp_path / "bad-lost-breakdown.xlsx",
            batch_name="DEMO-BAD-",
            invoice_date=date(2026, 8, 14),
        )


def test_no_clobber_is_atomic_and_extension_must_match_template(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    adapter: AccountingTemplateAdapter,
) -> None:
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"synthetic-winner")
    with pytest.raises(OutputExistsError):
        adapter.export(
            [_cases()[0]],
            existing,
            batch_name="DEMO-RACE-",
            invoice_date=date(2026, 8, 14),
        )
    assert existing.read_bytes() == b"synthetic-winner"

    with pytest.raises(AccountingValidationError, match="extension"):
        adapter.export(
            [_cases()[0]],
            tmp_path / "wrong.xlsm",
            batch_name="DEMO-RACE-",
            invoice_date=date(2026, 8, 14),
        )

    destination = tmp_path / "raced.xlsx"

    def create_competing_destination(source: object, target: object) -> None:
        del source
        Path(target).write_bytes(b"synthetic-race-winner")
        raise FileExistsError

    monkeypatch.setattr(
        "asset_compensation.adapters.accounting_template.os.link",
        create_competing_destination,
    )
    with pytest.raises(OutputExistsError):
        adapter.export(
            [_cases()[0]],
            destination,
            batch_name="DEMO-RACE-",
            invoice_date=date(2026, 8, 14),
        )
    assert destination.read_bytes() == b"synthetic-race-winner"


def test_xlsm_template_controls_output_suffix(tmp_path: Path) -> None:
    base = AccountingTemplateAdapter(
        {"DAMAGED": ("DEMO.DEBIT", "DEMO.CREDIT")}
    )
    macro_template = tmp_path / "synthetic-template.xlsm"
    shutil.copyfile(base.template_path, macro_template)
    adapter = AccountingTemplateAdapter(
        {"DAMAGED": ("DEMO.DEBIT", "DEMO.CREDIT")},
        template_path=macro_template,
    )

    result = adapter.export(
        [_cases()[0]],
        tmp_path / "synthetic-macro-output",
        batch_name="DEMO-MACRO-",
        invoice_date=date(2026, 8, 14),
    )

    assert adapter.template_path == macro_template.resolve()
    assert adapter.output_suffix == ".xlsm"
    assert result.path.suffix == ".xlsm"
    workbook = load_workbook(result.path, keep_vba=True)
    workbook.close()
