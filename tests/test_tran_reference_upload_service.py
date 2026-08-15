"""Safe storage tests for uploaded TranNNB reference workbooks."""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from asset_compensation.adapters import CcdcWorkbookIndex, FaGlWorkbookIndex
from asset_compensation.services.tran_reference_upload_service import (
    TranReferenceUpload,
    TranReferenceUploadError,
    TranReferenceUploadService,
)

_SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OFFICE_RELATIONSHIP_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_PACKAGE_RELATIONSHIP_NS = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_EXTERNAL_LINK_2021_NS = (
    "http://schemas.microsoft.com/office/spreadsheetml/2021/extlinks2021"
)
_EXTERNAL_LINK_RELATIONSHIP = f"{_OFFICE_RELATIONSHIP_NS}/externalLink"
_EXTERNAL_LINK_PATH_RELATIONSHIP = f"{_OFFICE_RELATIONSHIP_NS}/externalLinkPath"
_HYPERLINK_RELATIONSHIP = f"{_OFFICE_RELATIONSHIP_NS}/hyperlink"
_SHARED_STRINGS_RELATIONSHIP = f"{_OFFICE_RELATIONSHIP_NS}/sharedStrings"
_MAX_EXTERNAL_LINK_XML_BYTES = 1024 * 1024
_MAX_EXTERNAL_LINK_RELS_BYTES = 256 * 1024


def fa_gl_bytes(
    *,
    tag_number: str = "MOU10001",
    include_chart: bool = False,
    include_table: bool = False,
    data_validation_formula: str | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {
        "VNG-Asset": (3, 4),
        "VNG-Tool": (3, 4),
        "VNGS-Asset": (2, 3),
        "VNGS-Tool": (2, 3),
    }
    headers = {
        2: "Company Name",
        7: "Cost Center",
        8: "Product Code",
        10: "Location",
        12: "Asset",
        16: "Tag Number",
        22: "Date of Depreciation",
        25: "Life(month)",
        27: "Cost",
    }
    for sheet_name, (header_row, data_row) in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for column, header in headers.items():
            sheet.cell(header_row, column, header)
        if sheet_name == "VNG-Tool":
            sheet.cell(data_row, 2, "Synthetic Company")
            sheet.cell(data_row, 7, "0603")
            sheet.cell(data_row, 8, "000")
            sheet.cell(data_row, 10, "01")
            sheet.cell(data_row, 12, "SYN-001")
            sheet.cell(data_row, 16, tag_number)
            sheet.cell(data_row, 22, date(2025, 1, 1))
            sheet.cell(data_row, 25, 48)
            sheet.cell(data_row, 27, 1_000_000)
            if include_chart:
                chart = BarChart()
                chart.add_data(
                    Reference(sheet, min_col=27, min_row=header_row, max_row=data_row),
                    titles_from_data=True,
                )
                sheet.add_chart(chart, "AC6")
            if include_table:
                sheet["AC1"] = "Synthetic Input"
                sheet["AD1"] = "Synthetic Output"
                sheet["AC2"] = 1
                sheet["AD2"] = 2
                sheet.add_table(Table(displayName="SyntheticTable", ref="AC1:AD2"))
            if data_validation_formula:
                validation = DataValidation(
                    type="custom",
                    formula1=data_validation_formula,
                )
                sheet.add_data_validation(validation)
                validation.add("A1")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def ccdc_bytes() -> bytes:
    workbook = Workbook()
    define = workbook.active
    define.title = "Define"
    define.append(["Product Type", "Barcode", "Group Type"])
    define.append(["Other/Spe.part", "MOU", "Hardware"])
    warehouse = workbook.create_sheet("BC Xuatkho")
    warehouse.append(["Unused", "Asset Name", "Unused", "Unused", "Start time"])
    warehouse.append([None, "MOU10001", None, None, date(2025, 1, 1)])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def upload(data: bytes, filename: str) -> TranReferenceUpload:
    return TranReferenceUpload(
        filename=filename,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        stream=io.BytesIO(data),
    )


def with_external_book_metadata(
    data: bytes,
    *,
    target: str = "file:///synthetic/reference.xlsx",
    relationship_type: str = _EXTERNAL_LINK_PATH_RELATIONSHIP,
    link_kind: str = "externalBook",
    external_formula: str | None = None,
    external_defined_name: str | None = None,
    external_chart_formula: str | None = None,
    external_table_formula: str | None = None,
    ordinary_shared_string: str | None = None,
    cached_external_data: bool = False,
    include_rels: bool = True,
    extra_member: str | None = None,
    extra_xml_member: tuple[str, bytes] | None = None,
    chart_path_case_variant: bool = False,
    link_xml_size: int | None = None,
    rels_xml_size: int | None = None,
) -> bytes:
    """Add synthetic local-link metadata without borrowing an operational workbook."""

    with zipfile.ZipFile(io.BytesIO(data)) as source:
        payloads = {entry.filename: source.read(entry) for entry in source.infolist()}

    workbook = ET.fromstring(payloads["xl/workbook.xml"])
    external_references = ET.SubElement(
        workbook,
        f"{{{_SPREADSHEET_NS}}}externalReferences",
    )
    ET.SubElement(
        external_references,
        f"{{{_SPREADSHEET_NS}}}externalReference",
        {f"{{{_OFFICE_RELATIONSHIP_NS}}}id": "rIdSyntheticExternal"},
    )
    if external_defined_name:
        defined_names = ET.SubElement(workbook, f"{{{_SPREADSHEET_NS}}}definedNames")
        defined_name = ET.SubElement(
            defined_names,
            f"{{{_SPREADSHEET_NS}}}definedName",
            {"name": "SyntheticExternalName"},
        )
        defined_name.text = external_defined_name
    payloads["xl/workbook.xml"] = ET.tostring(
        workbook,
        encoding="utf-8",
        xml_declaration=True,
    )

    workbook_rels = ET.fromstring(payloads["xl/_rels/workbook.xml.rels"])
    ET.SubElement(
        workbook_rels,
        f"{{{_PACKAGE_RELATIONSHIP_NS}}}Relationship",
        {
            "Id": "rIdSyntheticExternal",
            "Type": _EXTERNAL_LINK_RELATIONSHIP,
            "Target": "externalLinks/externalLink1.xml",
        },
    )
    payloads["xl/_rels/workbook.xml.rels"] = ET.tostring(
        workbook_rels,
        encoding="utf-8",
        xml_declaration=True,
    )

    content_types = ET.fromstring(payloads["[Content_Types].xml"])
    ET.SubElement(
        content_types,
        f"{{{_CONTENT_TYPES_NS}}}Override",
        {
            "PartName": "/xl/externalLinks/externalLink1.xml",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.externalLink+xml"
            ),
        },
    )
    payloads["[Content_Types].xml"] = ET.tostring(
        content_types,
        encoding="utf-8",
        xml_declaration=True,
    )

    external_link = ET.Element(f"{{{_SPREADSHEET_NS}}}externalLink")
    if link_kind == "externalBook":
        external_book = ET.SubElement(
            external_link,
            f"{{{_SPREADSHEET_NS}}}externalBook",
            {f"{{{_OFFICE_RELATIONSHIP_NS}}}id": "rIdSource"},
        )
        alternate_urls = ET.SubElement(
            external_book,
            f"{{{_EXTERNAL_LINK_2021_NS}}}alternateUrls",
        )
        ET.SubElement(
            alternate_urls,
            f"{{{_EXTERNAL_LINK_2021_NS}}}absoluteUrl",
            {f"{{{_OFFICE_RELATIONSHIP_NS}}}id": "rIdAlternate"},
        )
        sheet_names = ET.SubElement(
            external_book,
            f"{{{_SPREADSHEET_NS}}}sheetNames",
        )
        ET.SubElement(
            sheet_names,
            f"{{{_SPREADSHEET_NS}}}sheetName",
            {"val": "Synthetic"},
        )
        sheet_data_set = ET.SubElement(
            external_book,
            f"{{{_SPREADSHEET_NS}}}sheetDataSet",
        )
        sheet_data = ET.SubElement(
            sheet_data_set,
            f"{{{_SPREADSHEET_NS}}}sheetData",
            {"sheetId": "0"},
        )
        if cached_external_data:
            ET.SubElement(sheet_data, f"{{{_SPREADSHEET_NS}}}row", {"r": "1"})
    else:
        ET.SubElement(external_link, f"{{{_SPREADSHEET_NS}}}{link_kind}")
    payloads["xl/externalLinks/externalLink1.xml"] = ET.tostring(
        external_link,
        encoding="utf-8",
        xml_declaration=True,
    )

    if include_rels:
        relationships = ET.Element(f"{{{_PACKAGE_RELATIONSHIP_NS}}}Relationships")
        for relationship_id in ("rIdSource", "rIdAlternate"):
            ET.SubElement(
                relationships,
                f"{{{_PACKAGE_RELATIONSHIP_NS}}}Relationship",
                {
                    "Id": relationship_id,
                    "Type": relationship_type,
                    "Target": target,
                    "TargetMode": "External",
                },
            )
        payloads["xl/externalLinks/_rels/externalLink1.xml.rels"] = ET.tostring(
            relationships,
            encoding="utf-8",
            xml_declaration=True,
        )

    if external_formula:
        worksheet = ET.fromstring(payloads["xl/worksheets/sheet1.xml"])
        sheet_data = worksheet.find(f"{{{_SPREADSHEET_NS}}}sheetData")
        assert sheet_data is not None
        row = ET.SubElement(sheet_data, f"{{{_SPREADSHEET_NS}}}row", {"r": "1"})
        cell = ET.SubElement(row, f"{{{_SPREADSHEET_NS}}}c", {"r": "A1"})
        formula = ET.SubElement(cell, f"{{{_SPREADSHEET_NS}}}f")
        formula.text = external_formula
        ET.SubElement(cell, f"{{{_SPREADSHEET_NS}}}v").text = "0"
        payloads["xl/worksheets/sheet1.xml"] = ET.tostring(
            worksheet,
            encoding="utf-8",
            xml_declaration=True,
        )

    if external_chart_formula:
        chart_name = next(
            name
            for name in payloads
            if name.casefold().startswith("xl/charts/")
            and name.casefold().endswith(".xml")
        )
        chart = ET.fromstring(payloads[chart_name])
        chart_formula = next(
            element
            for element in chart.iter()
            if str(element.tag).rsplit("}", 1)[-1] == "f"
        )
        chart_formula.text = external_chart_formula
        payloads[chart_name] = ET.tostring(
            chart,
            encoding="utf-8",
            xml_declaration=True,
        )
        if chart_path_case_variant:
            variant_name = chart_name.replace("xl/charts/", "xl/Charts/")
            payloads[variant_name] = payloads.pop(chart_name)
            for name in tuple(payloads):
                if not name.casefold().endswith(".rels"):
                    continue
                relationships = ET.fromstring(payloads[name])
                changed = False
                for relationship in relationships:
                    target_value = relationship.attrib.get("Target", "")
                    replacement = target_value.replace(
                        "charts/chart1.xml",
                        "Charts/chart1.xml",
                    )
                    if replacement != target_value:
                        relationship.attrib["Target"] = replacement
                        changed = True
                if changed:
                    payloads[name] = ET.tostring(
                        relationships,
                        encoding="utf-8",
                        xml_declaration=True,
                    )
            content_types = ET.fromstring(payloads["[Content_Types].xml"])
            for override in content_types:
                part_name = override.attrib.get("PartName", "")
                replacement = part_name.replace(
                    "/xl/charts/chart1.xml",
                    "/xl/Charts/chart1.xml",
                )
                if replacement != part_name:
                    override.attrib["PartName"] = replacement
            payloads["[Content_Types].xml"] = ET.tostring(
                content_types,
                encoding="utf-8",
                xml_declaration=True,
            )

    if external_table_formula:
        table_name = next(
            name
            for name in payloads
            if name.casefold().startswith("xl/tables/")
            and name.casefold().endswith(".xml")
        )
        table = ET.fromstring(payloads[table_name])
        table_columns = [
            element
            for element in table.iter()
            if str(element.tag).rsplit("}", 1)[-1] == "tableColumn"
        ]
        calculated_formula = ET.SubElement(
            table_columns[-1],
            f"{{{_SPREADSHEET_NS}}}calculatedColumnFormula",
        )
        calculated_formula.text = external_table_formula
        payloads[table_name] = ET.tostring(
            table,
            encoding="utf-8",
            xml_declaration=True,
        )

    if ordinary_shared_string:
        workbook_rels = ET.fromstring(payloads["xl/_rels/workbook.xml.rels"])
        ET.SubElement(
            workbook_rels,
            f"{{{_PACKAGE_RELATIONSHIP_NS}}}Relationship",
            {
                "Id": "rIdSyntheticSharedStrings",
                "Type": _SHARED_STRINGS_RELATIONSHIP,
                "Target": "sharedStrings.xml",
            },
        )
        payloads["xl/_rels/workbook.xml.rels"] = ET.tostring(
            workbook_rels,
            encoding="utf-8",
            xml_declaration=True,
        )
        content_types = ET.fromstring(payloads["[Content_Types].xml"])
        ET.SubElement(
            content_types,
            f"{{{_CONTENT_TYPES_NS}}}Override",
            {
                "PartName": "/xl/sharedStrings.xml",
                "ContentType": (
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sharedStrings+xml"
                ),
            },
        )
        payloads["[Content_Types].xml"] = ET.tostring(
            content_types,
            encoding="utf-8",
            xml_declaration=True,
        )
        shared_strings = ET.Element(
            f"{{{_SPREADSHEET_NS}}}sst",
            {"count": "1", "uniqueCount": "1"},
        )
        shared_item = ET.SubElement(shared_strings, f"{{{_SPREADSHEET_NS}}}si")
        ET.SubElement(shared_item, f"{{{_SPREADSHEET_NS}}}t").text = ordinary_shared_string
        payloads["xl/sharedStrings.xml"] = ET.tostring(
            shared_strings,
            encoding="utf-8",
            xml_declaration=True,
        )
        worksheet = ET.fromstring(payloads["xl/worksheets/sheet1.xml"])
        sheet_data = worksheet.find(f"{{{_SPREADSHEET_NS}}}sheetData")
        assert sheet_data is not None
        row = ET.SubElement(sheet_data, f"{{{_SPREADSHEET_NS}}}row", {"r": "1"})
        cell = ET.SubElement(
            row,
            f"{{{_SPREADSHEET_NS}}}c",
            {"r": "A1", "t": "s"},
        )
        ET.SubElement(cell, f"{{{_SPREADSHEET_NS}}}v").text = "0"
        payloads["xl/worksheets/sheet1.xml"] = ET.tostring(
            worksheet,
            encoding="utf-8",
            xml_declaration=True,
        )

    if extra_member:
        payloads[extra_member] = b"synthetic unsupported content"
    if extra_xml_member:
        payloads[extra_xml_member[0]] = extra_xml_member[1]

    stored_members: set[str] = set()
    for name, requested_size in (
        ("xl/externalLinks/externalLink1.xml", link_xml_size),
        ("xl/externalLinks/_rels/externalLink1.xml.rels", rels_xml_size),
    ):
        if requested_size is None:
            continue
        assert len(payloads[name]) <= requested_size
        payloads[name] += b" " * (requested_size - len(payloads[name]))
        stored_members.add(name)

    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in payloads.items():
            archive.writestr(
                name,
                payload,
                compress_type=(
                    zipfile.ZIP_STORED if name in stored_members else zipfile.ZIP_DEFLATED
                ),
            )
    return output.getvalue()


def test_upload_activates_atomic_pair_preserves_optional_ccdc_and_clears_explicitly(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")

    first = service.upload(
        upload(fa_gl_bytes(), "client-fa.xlsx"),
        upload(ccdc_bytes(), "client-ccdc.xlsx"),
    )
    first_snapshot = service.snapshot()
    second = service.upload(upload(fa_gl_bytes(tag_number="MOU10002"), "new-fa.xlsx"))
    second_snapshot = service.snapshot()

    assert first == {
        "configured": True,
        "updated_at": first["updated_at"],
        "fa_gl_configured": True,
        "ccdc_configured": True,
    }
    assert first_snapshot.fa_gl_path is not None
    assert first_snapshot.ccdc_path is not None
    assert second["ccdc_configured"] is True
    assert second_snapshot.ccdc_path is not None
    assert "client-fa" not in str(service.status())
    assert str(tmp_path) not in str(service.status())

    cleared_ccdc = service.upload(
        upload(fa_gl_bytes(), "fa.xlsx"),
        clear_ccdc=True,
    )
    assert cleared_ccdc["ccdc_configured"] is False
    assert service.snapshot().ccdc_path is None


def test_invalid_replacement_does_not_change_active_version(tmp_path: Path) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    service.upload(upload(fa_gl_bytes(), "fa.xlsx"))
    before = service.snapshot().fa_gl_path

    with pytest.raises(TranReferenceUploadError):
        service.upload(upload(b"not an xlsx", "replacement.xlsx"))

    assert service.snapshot().fa_gl_path == before


def test_xlsx_preflight_allows_inert_local_external_book_metadata(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")

    result = service.upload(
        upload(with_external_book_metadata(fa_gl_bytes()), "fa.xlsx"),
        upload(with_external_book_metadata(ccdc_bytes()), "ccdc.xlsx"),
    )
    snapshot = service.snapshot()

    assert result["fa_gl_configured"] is True
    assert result["ccdc_configured"] is True
    assert snapshot.fa_gl_path is not None
    assert snapshot.ccdc_path is not None
    assert FaGlWorkbookIndex.from_path(snapshot.fa_gl_path).lookup("MOU10001").record
    assert CcdcWorkbookIndex.from_path(snapshot.ccdc_path).classification("MOU").matches


def test_xlsx_preflight_allows_external_token_as_ordinary_shared_string(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(
        fa_gl_bytes(),
        ordinary_shared_string="[Book.xlsx] is documentation text, not a formula",
    )

    result = service.upload(upload(data, "fa.xlsx"))

    assert result["fa_gl_configured"] is True


@pytest.mark.parametrize(
    "formula",
    [
        "SUM(SyntheticTable[2025])",
        "SUM([@2025])",
        '="[1]"',
        "SUM(Table1[Column])",
        "SUM([[#Headers],[Book.csv]])",
        "SUM(SyntheticTable[#Headers])",
        "SUM(SyntheticTable[[#Headers],[Book.xlsx]])",
        "SUM([Book.csv]:[Template.xltx])",
        "SUM([Book.csv] [Template.xltx])",
        "SUM([Col] Table2[Other])",
    ],
)
def test_xlsx_preflight_allows_internal_brackets_in_formula_contexts(
    tmp_path: Path,
    formula: str,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(
        fa_gl_bytes(),
        external_formula=formula,
    )

    result = service.upload(upload(data, "fa.xlsx"))

    assert result["fa_gl_configured"] is True


@pytest.mark.parametrize(
    "member",
    [
        "xl/vbaProject.bin",
        "xl/embeddings/payload.bin",
        "xl/activeX/activeX1.bin",
        "xl/oleObjects/oleObject1.bin",
    ],
)
def test_xlsx_preflight_rejects_active_or_embedded_content_before_activation(
    tmp_path: Path,
    member: str,
) -> None:
    malicious = io.BytesIO(fa_gl_bytes())
    with zipfile.ZipFile(malicious, mode="a") as archive:
        archive.writestr(member, b"synthetic payload")
    service = TranReferenceUploadService(tmp_path / "references")

    with pytest.raises(TranReferenceUploadError, match="active, linked, or embedded"):
        service.upload(upload(malicious.getvalue(), "fa.xlsx"))

    assert service.status()["configured"] is False


@pytest.mark.parametrize("link_kind", ["ddeLink", "oleLink"])
def test_xlsx_preflight_rejects_active_external_link_kinds(
    tmp_path: Path,
    link_kind: str,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(fa_gl_bytes(), link_kind=link_kind)

    with pytest.raises(TranReferenceUploadError, match="unused local-file metadata"):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


@pytest.mark.parametrize(
    ("target", "relationship_type"),
    [
        ("https://example.invalid/reference.xlsx", _EXTERNAL_LINK_PATH_RELATIONSHIP),
        ("file://[invalid", _EXTERNAL_LINK_PATH_RELATIONSHIP),
        ("file:///synthetic/reference.xlsx", _HYPERLINK_RELATIONSHIP),
    ],
)
def test_xlsx_preflight_rejects_nonlocal_or_wrong_external_link_relationship(
    tmp_path: Path,
    target: str,
    relationship_type: str,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(
        fa_gl_bytes(),
        target=target,
        relationship_type=relationship_type,
    )

    with pytest.raises(TranReferenceUploadError, match="unused local-file metadata"):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


@pytest.mark.parametrize(
    "options",
    [
        {"external_formula": "[1]Synthetic!A1"},
        {"external_formula": "@[1]Synthetic!A1"},
        {"external_formula": "[Book.xlsx]Synthetic!A1"},
        {"external_formula": "[Book.csv]Synthetic!A1"},
        {"external_formula": "LocalRange [Book.csv]Synthetic!A1"},
        {"external_formula": "A1 [Book.csv]Sheet!A1"},
        {"external_formula": "[Local] [Book.csv]Synthetic!A1"},
        {"external_formula": "[#Book.csv]Synthetic!A1"},
        {"external_formula": "[@Book.csv]Synthetic!A1"},
        {"external_formula": "'[Book[Draft].csv]Synthetic'!A1"},
        {"external_formula": "'[Template.xltx]Synthetic'!A1"},
        {"external_formula": "'[Book.xlsx]Synthetic'!A1"},
        {"external_defined_name": "[Addin.xlam]ExternalName"},
        {"external_defined_name": "[Book.xlsx]ExternalName"},
        {"external_defined_name": "[Book.xlsm]Synthetic!$A$1"},
    ],
)
def test_xlsx_preflight_rejects_external_formula_or_defined_name(
    tmp_path: Path,
    options: dict[str, str],
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(fa_gl_bytes(), **options)

    with pytest.raises(
        TranReferenceUploadError,
        match="external workbook formula or defined name",
    ):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


@pytest.mark.parametrize(
    ("kind", "expected_part"),
    [
        ("chart", "xl/Charts/chart1.xml"),
        ("table", "xl/tables/table1.xml"),
        ("data-validation", "xl/worksheets/sheet2.xml"),
    ],
)
def test_xlsx_preflight_rejects_external_formula_in_related_ooxml_part(
    tmp_path: Path,
    kind: str,
    expected_part: str,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    if kind == "chart":
        workbook_data = fa_gl_bytes(include_chart=True)
        options: dict[str, str | bool] = {
            "external_chart_formula": "[1]Synthetic!$A$1:$A$2",
            "chart_path_case_variant": True,
        }
    elif kind == "table":
        workbook_data = fa_gl_bytes(include_table=True)
        options = {"external_table_formula": "[Book.xlsx]Synthetic!A1"}
    else:
        workbook_data = fa_gl_bytes(
            data_validation_formula="[Book.xlsb]Synthetic!A1"
        )
        options = {}
    data = with_external_book_metadata(workbook_data, **options)

    with pytest.raises(
        TranReferenceUploadError,
        match="external workbook formula or defined name",
    ) as caught:
        service.upload(upload(data, "fa.xlsx"))

    assert expected_part in str(caught.value)
    assert service.status()["configured"] is False


def test_xlsx_preflight_rejects_external_formula_without_link_metadata(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = fa_gl_bytes(data_validation_formula="[Book.xlsx]Synthetic!A1")

    with pytest.raises(
        TranReferenceUploadError,
        match="external workbook formula or defined name",
    ):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


def test_xlsx_preflight_rejects_external_formula_attribute(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    formula_part = (
        f'<pivotCacheDefinition xmlns="{_SPREADSHEET_NS}">'
        '<calculatedItems><calculatedItem formula="[1]Synthetic!A1" />'
        "</calculatedItems></pivotCacheDefinition>"
    ).encode()
    data = with_external_book_metadata(
        fa_gl_bytes(),
        extra_xml_member=(
            "xl/pivotCache/pivotCacheDefinition1.xml",
            formula_part,
        ),
    )

    with pytest.raises(
        TranReferenceUploadError,
        match="external workbook formula or defined name",
    ) as caught:
        service.upload(upload(data, "fa.xlsx"))

    assert "xl/pivotCache/pivotCacheDefinition1.xml" in str(caught.value)
    assert service.status()["configured"] is False


def test_xlsx_preflight_rejects_deep_unrelated_xml_part(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    deep_xml = b"<root>" + (b"<node>" * 65) + (b"</node>" * 65) + b"</root>"
    data = with_external_book_metadata(
        fa_gl_bytes(),
        extra_xml_member=("xl/deep.xml", deep_xml),
    )

    with pytest.raises(TranReferenceUploadError, match="safe XML complexity limit"):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


def test_xlsx_preflight_accepts_external_link_xml_at_exact_size_limits(
    tmp_path: Path,
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(
        fa_gl_bytes(),
        link_xml_size=_MAX_EXTERNAL_LINK_XML_BYTES,
        rels_xml_size=_MAX_EXTERNAL_LINK_RELS_BYTES,
    )

    result = service.upload(upload(data, "fa.xlsx"))

    assert result["fa_gl_configured"] is True


@pytest.mark.parametrize(
    "options",
    [
        {"link_xml_size": _MAX_EXTERNAL_LINK_XML_BYTES + 1},
        {"rels_xml_size": _MAX_EXTERNAL_LINK_RELS_BYTES + 1},
    ],
)
def test_xlsx_preflight_rejects_external_link_xml_over_size_limit(
    tmp_path: Path,
    options: dict[str, int],
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(fa_gl_bytes(), **options)

    with pytest.raises(TranReferenceUploadError, match="safe per-link size limit"):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


@pytest.mark.parametrize(
    "options",
    [
        {"cached_external_data": True},
        {"include_rels": False},
        {"extra_member": "xl/externalLinks/unexpected.bin"},
    ],
)
def test_xlsx_preflight_rejects_non_orphan_or_unpaired_external_metadata(
    tmp_path: Path,
    options: dict[str, bool | str],
) -> None:
    service = TranReferenceUploadService(tmp_path / "references")
    data = with_external_book_metadata(fa_gl_bytes(), **options)

    with pytest.raises(TranReferenceUploadError, match="unused local-file metadata"):
        service.upload(upload(data, "fa.xlsx"))

    assert service.status()["configured"] is False


def test_xlsx_preflight_rejects_external_relationship_before_activation(
    tmp_path: Path,
) -> None:
    malicious = io.BytesIO(fa_gl_bytes())
    with zipfile.ZipFile(malicious, mode="a") as archive:
        archive.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            b"""<?xml version="1.0" encoding="UTF-8"?>
            <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
              <Relationship Id="rId99" Type="synthetic" Target="https://example.invalid/"
                            TargetMode="External" />
            </Relationships>""",
        )
    service = TranReferenceUploadService(tmp_path / "references")

    with pytest.raises(TranReferenceUploadError, match="external relationship"):
        service.upload(upload(malicious.getvalue(), "fa.xlsx"))

    assert service.status()["configured"] is False


@pytest.mark.parametrize(
    "relationship_xml",
    [
        b"<Relationships>"
        + (b"<node>" * 16)
        + (b"</node>" * 16)
        + b"</Relationships>",
        (
            b'<!DOCTYPE Relationships [<!ENTITY synthetic "payload">]>'
            b"<Relationships>&synthetic;</Relationships>"
        ),
    ],
    ids=["deep", "dtd"],
)
def test_xlsx_preflight_rejects_unsafe_relationship_xml(
    tmp_path: Path,
    relationship_xml: bytes,
) -> None:
    malicious = io.BytesIO(fa_gl_bytes())
    with zipfile.ZipFile(malicious, mode="a") as archive:
        archive.writestr("xl/_rels/synthetic.xml.rels", relationship_xml)
    service = TranReferenceUploadService(tmp_path / "references")

    with pytest.raises(
        TranReferenceUploadError,
        match="safe relationship XML complexity limit",
    ):
        service.upload(upload(malicious.getvalue(), "fa.xlsx"))

    assert service.status()["configured"] is False


def test_clear_removes_only_managed_versions(tmp_path: Path) -> None:
    reference_root = tmp_path / "references"
    service = TranReferenceUploadService(reference_root)
    service.upload(upload(fa_gl_bytes(), "fa.xlsx"))
    sentinel = reference_root / "keep-me.txt"
    sentinel.write_text("unknown file remains", "utf-8")

    counts = service.clear()

    assert counts["tran_reference_version_count"] == 1
    assert counts["tran_fa_gl_reference_count"] == 1
    assert service.status()["configured"] is False
    assert sentinel.read_text("utf-8") == "unknown file remains"
