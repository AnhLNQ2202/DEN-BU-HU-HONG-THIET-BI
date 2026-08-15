"""End-to-end HTTP contracts for the safe TranNNB workflow."""

from __future__ import annotations

import io
from datetime import date
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path
from threading import Event, Thread

from flask import Flask
from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

import asset_compensation.web.routes as routes_module
from asset_compensation.config import Settings
from asset_compensation.services import MailPdfService
from asset_compensation.web import create_app


class _SyntheticPdfConverter:
    def convert_eml(
        self,
        source: str | Path,
        destination: str | Path,
        *,
        overwrite: bool = False,
    ) -> Path:
        assert Path(source).is_file()
        output = Path(destination)
        assert overwrite is False
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with output.open("xb") as stream:
            writer.write(stream)
        return output


def _fa_gl_bytes() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {
        "VNG-Asset": (3, 4),
        "VNG-Tool": (3, 4),
        "VNGS-Asset": (2, 3),
        "VNGS-Tool": (2, 3),
    }
    headers = {
        2: "Company Name",
        7: "Cost Center",
        8: "Product Code",
        10: "Location",
        12: "Asset",
        16: "Tag Number",
        22: "Date of Depreciation",
        25: "Life(month)",
        27: "Cost",
    }
    for sheet_name, (header_row, data_row) in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for column, header in headers.items():
            sheet.cell(header_row, column, header)
        if sheet_name == "VNG-Tool":
            sheet.cell(data_row, 2, "Synthetic Company")
            sheet.cell(data_row, 7, "0603")
            sheet.cell(data_row, 8, "000")
            sheet.cell(data_row, 10, "01")
            sheet.cell(data_row, 12, "SYN-001")
            sheet.cell(data_row, 16, "MOU10001")
            sheet.cell(data_row, 22, date(2025, 1, 1))
            sheet.cell(data_row, 25, 48)
            sheet.cell(data_row, 27, 1_000_000)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _lost_table_email() -> bytes:
    message = EmailMessage()
    message["Subject"] = "IT - Thông tin tài sản thất lạc"
    message["From"] = "Synthetic Asset Team <asset@example.invalid>"
    message["To"] = (
        "Operator <operator@example.invalid>, Manager <manager@example.invalid>"
    )
    message["Cc"] = "Audit <audit@example.invalid>"
    message["Date"] = "Thu, 13 Aug 2026 03:49:56 +0000"
    message["Message-ID"] = "<synthetic-retained@example.invalid>"
    message.set_content("Synthetic lost-asset table follows.")
    message.add_alternative(
        """
        <html><body><table>
          <tr>
            <th>Asset Name</th><th>Product Name</th><th>Domain</th>
            <th>Ngày bắt đầu sử dụng</th><th>Ngày thất lạc/mất</th>
            <th>Nguyên giá ban đầu</th><th>Mức khấu hao sử dụng còn lại</th>
            <th>Phí đền bù trách nhiệm</th><th>Tổng số tiền đền bù</th>
            <th>NOTE</th><th>Entity</th><th>Cost center</th>
            <th>Product code</th><th>Location</th>
          </tr>
          <tr>
            <td>DEMO-LAP-101</td><td>Synthetic Laptop</td><td>demo.alpha</td>
            <td>01/01/2025</td><td>01/08/2026</td><td>1,000</td>
            <td>600</td><td>100</td><td>700</td><td>Asset</td><td>VNG</td>
            <td>'0603</td><td>'000</td><td>'01</td>
          </tr>
          <tr>
            <td>DEMO-MON-102</td><td>Synthetic Monitor</td><td>demo.beta</td>
            <td>01/01/2025</td><td>01/08/2026</td><td>500</td>
            <td>200</td><td>50</td><td>250</td><td>Asset</td><td>VNG</td>
            <td>'0603</td><td>'000</td><td>'01</td>
          </tr>
        </table></body></html>
        """,
        subtype="html",
    )
    return message.as_bytes()


def _asset() -> dict[str, object]:
    return {
        "tag_number": "MOU10001",
        "asset_name": "Synthetic mouse",
        "domain": "demo.user",
        "lost_date": "2026-01-01",
    }


def _app(tmp_path: Path, **overrides: object) -> Flask:
    settings = {
        "data_dir": tmp_path / "runtime",
        "demo_mode": False,
        "secret_key": "synthetic-tran-api-key",
    }
    settings.update(overrides)
    app = create_app(Settings(**settings))
    app.config.update(TESTING=True)
    return app


def _upload_fa(app: Flask) -> dict[str, object]:
    response = app.test_client().post(
        "/api/tran/references/upload",
        data={
            "fa_gl_file": (
                io.BytesIO(_fa_gl_bytes()),
                "synthetic-fa-gl.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={"X-Asset-Hub-Upload": "tran-reference-v1"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 201
    return response.get_json()


def test_capabilities_fail_closed_until_reference_is_uploaded(tmp_path: Path) -> None:
    app = _app(tmp_path)
    try:
        before = app.test_client().get("/api/capabilities").get_json()
        unresolved = app.test_client().post("/api/tran/resolve", json={"assets": [_asset()]})
        uploaded = _upload_fa(app)
        after = app.test_client().get("/api/capabilities").get_json()

        assert before["capabilities"]["raw_eml_retention"] is False
        assert before["capabilities"]["mail_pdf_individual"] is False
        assert before["tran_references"]["fa_gl"]["available"] is False
        assert unresolved.status_code == 400
        assert uploaded["status"]["fa_gl"] == {
            "configured": True,
            "available": True,
            "source": "uploaded",
        }
        assert str(tmp_path) not in str(uploaded)
        assert "synthetic-fa-gl" not in str(uploaded)
        assert after["capabilities"]["tran_lookup"] is True
        assert after["capabilities"]["tran_workbook_export"] is True
    finally:
        app.extensions["asset_hub"]["repository"].close()


def test_tran_resolve_serializes_against_reference_replacement(
    tmp_path: Path,
    monkeypatch,
) -> None:
    app = _app(tmp_path)
    _upload_fa(app)
    resolve_started = Event()
    release_resolve = Event()
    upload_started = Event()
    upload_finished = Event()
    failures: list[BaseException] = []
    responses: dict[str, int] = {}
    original_from_path = routes_module.FaGlWorkbookIndex.from_path

    def paused_from_path(path: str | Path) -> object:
        resolve_started.set()
        if not release_resolve.wait(timeout=5):
            raise RuntimeError("Synthetic resolve was not released")
        return original_from_path(path)

    def run_resolve() -> None:
        try:
            with app.test_client() as client:
                responses["resolve"] = client.post(
                    "/api/tran/resolve",
                    json={"assets": [_asset()]},
                ).status_code
        except BaseException as exc:  # pragma: no cover - surfaced below
            failures.append(exc)

    def run_upload() -> None:
        upload_started.set()
        try:
            _upload_fa(app)
            responses["upload"] = 201
        except BaseException as exc:  # pragma: no cover - surfaced below
            failures.append(exc)
        finally:
            upload_finished.set()

    monkeypatch.setattr(routes_module.FaGlWorkbookIndex, "from_path", paused_from_path)
    resolve_thread = Thread(target=run_resolve)
    upload_thread = Thread(target=run_upload)
    try:
        resolve_thread.start()
        assert resolve_started.wait(timeout=5)
        upload_thread.start()
        assert upload_started.wait(timeout=5)
        assert not upload_finished.wait(timeout=0.1)
        release_resolve.set()
        resolve_thread.join(timeout=5)
        upload_thread.join(timeout=5)

        assert not failures
        assert not resolve_thread.is_alive()
        assert not upload_thread.is_alive()
        assert responses == {"resolve": 200, "upload": 201}
    finally:
        release_resolve.set()
        resolve_thread.join(timeout=5)
        upload_thread.join(timeout=5)
        app.extensions["asset_hub"]["repository"].close()


def test_resolve_export_and_download_use_bundled_clean_template(tmp_path: Path) -> None:
    app = _app(tmp_path)
    try:
        _upload_fa(app)
        client = app.test_client()
        resolved = client.post("/api/tran/resolve", json={"assets": [_asset()]})
        exported = client.post(
            "/api/tran/workbooks",
            json={"assets": [_asset()], "processing_date": "2026-01-15"},
        )

        assert resolved.status_code == 200
        assert resolved.get_json()["ready"] is True
        assert "MOU10001" in resolved.get_json()["mail_table_html"]
        assert exported.status_code == 201
        download = client.get(exported.get_json()["download_url"])
        assert download.status_code == 200
        assert download.headers["Cache-Control"] == "private, no-store"

        workbook = load_workbook(io.BytesIO(download.data), data_only=False)
        try:
            assert "2026" in workbook.sheetnames
            assert workbook["2026"]["D4"].value == "MOU10001"
            assert workbook["Sent out"]["A2"].value == "MOU10001"
        finally:
            workbook.close()
        assert client.post("/api/tran/drafts", json={}).status_code == 503
        assert client.post("/api/mail-pdfs/individual", json={}).status_code == 503
    finally:
        app.extensions["asset_hub"]["repository"].close()


def test_retained_multi_case_email_is_downloadable_and_builds_unsent_draft(
    tmp_path: Path,
) -> None:
    fa_path = tmp_path / "external-fa.xlsx"
    fa_bytes = _fa_gl_bytes()
    fa_path.write_bytes(fa_bytes)
    app = _app(
        tmp_path,
        retain_raw_eml=True,
        fa_gl_reference=fa_path,
        draft_from_address="operator@example.invalid",
        allow_test_reset=True,
    )
    app.extensions["asset_hub"]["mail_pdf_available"] = False
    try:
        client = app.test_client()
        source_email = _lost_table_email()
        uploaded = client.post(
            "/api/emails/upload",
            data={
                "files": (
                    io.BytesIO(source_email),
                    "Synthetic Evidence.eml",
                    "message/rfc822",
                )
            },
            headers={"X-Asset-Hub-Upload": "email-v1"},
            content_type="multipart/form-data",
        )

        assert uploaded.status_code == 200
        payload = uploaded.get_json()
        assert payload["ingested"] == 2
        assert payload["message"] == "Created 2 cases from 1 uploaded email"
        assert payload["retained_source_count"] == 1
        handles = {case["source_eml"]["handle"] for case in payload["cases"]}
        filenames = {case["source_eml"]["filename"] for case in payload["cases"]}
        assert len(handles) == 1
        assert filenames == {"Synthetic_Evidence.eml"}
        dashboard_cases = client.get("/api/dashboard").get_json()["cases"]
        assert {
            item["source_eml"]["download_url"] for item in dashboard_cases
        } == {payload["cases"][0]["source_eml"]["download_url"]}
        source = client.get(payload["cases"][0]["source_eml"]["download_url"])
        assert source.status_code == 200
        assert source.data == source_email
        assert source.headers["Cache-Control"] == "private, no-store"
        source.close()

        handle = handles.pop()
        draft = client.post(
            "/api/tran/drafts",
            json={
                "assets": [_asset()],
                "mail_artifact_handle": handle,
                "body_intro": "Synthetic approved response",
                "processing_date": "2026-01-15",
            },
        )
        assert draft.status_code == 201
        assert draft.get_json()["sent"] is False
        downloaded = client.get(draft.get_json()["draft_download_url"])
        assert downloaded.headers["Cache-Control"] == "private, no-store"
        parsed = BytesParser(policy=policy.default).parsebytes(downloaded.data)
        assert parsed["X-Unsent"] == "1"
        assert "asset@example.invalid" in str(parsed["To"])
        assert "operator@example.invalid" not in str(parsed["To"])
        downloaded.close()

        cleared = client.post(
            "/api/test-data/clear",
            json={"confirm": "CLEAR_TEST_DATA"},
            headers={"X-Asset-Hub-Action": "clear-test-data-v1"},
        )
        assert cleared.status_code == 200
        assert cleared.get_json()["mail_artifact_count"] == 1
        assert cleared.get_json()["tran_managed_output_count"] == 2
        assert client.get(f"/api/mail-artifacts/{handle}/download").status_code == 400
        assert client.get(draft.get_json()["draft_download_url"]).status_code == 400
        assert fa_path.read_bytes() == fa_bytes
    finally:
        app.extensions["asset_hub"]["repository"].close()


def test_settings_raw_retention_is_explicit_opt_in(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setenv("ASSET_HUB_DATA_DIR", str(tmp_path / "runtime"))
    assert Settings.from_env().retain_raw_eml is False
    assert not (tmp_path / "runtime" / "private-mail-artifacts").exists()

    monkeypatch.setenv("ASSET_HUB_RETAIN_RAW_EML", "true")
    settings = Settings.from_env()
    assert settings.retain_raw_eml is True
    assert settings.mail_artifact_dir == tmp_path / "runtime" / "private-mail-artifacts"


def test_unclassified_uploaded_email_is_not_retained(tmp_path: Path) -> None:
    app = _app(tmp_path, retain_raw_eml=True)
    message = EmailMessage()
    message["Subject"] = "Synthetic unclassified note"
    message["From"] = "sender@example.invalid"
    message["To"] = "recipient@example.invalid"
    message["Message-ID"] = "<unclassified@example.invalid>"
    message.set_content("No compensation record is present in this synthetic email.")
    try:
        response = app.test_client().post(
            "/api/emails/upload",
            data={
                "files": (
                    io.BytesIO(message.as_bytes()),
                    "unclassified.eml",
                    "message/rfc822",
                )
            },
            headers={"X-Asset-Hub-Upload": "email-v1"},
            content_type="multipart/form-data",
        )

        assert response.status_code == 200
        payload = response.get_json()
        assert payload["ingested"] == 0
        assert payload["retained_source_count"] == 0
        assert payload["unknown_files"] == ["upload-01.eml"]
        assert not list(app.extensions["asset_hub"]["settings"].mail_artifact_dir.rglob("*.eml"))
    finally:
        app.extensions["asset_hub"]["repository"].close()


def test_cloud_pdf_endpoints_create_individual_and_merged_downloads(
    tmp_path: Path,
) -> None:
    app = _app(tmp_path, retain_raw_eml=True, allow_test_reset=True)
    extension = app.extensions["asset_hub"]
    extension["mail_pdf_service"] = MailPdfService(
        extension["mail_artifact_store"],
        _SyntheticPdfConverter(),
    )
    extension["mail_pdf_available"] = True
    extension["pypdf_available"] = True
    extension["pdf_backend"] = "synthetic-test"
    try:
        client = app.test_client()
        capabilities = client.get("/api/capabilities").get_json()["capabilities"]
        assert capabilities["mail_pdf_individual"] is True
        assert capabilities["mail_pdf_batch"] is True
        assert capabilities["mail_pdf_backend"] == "synthetic-test"
        uploaded = client.post(
            "/api/emails/upload",
            data={
                "files": (
                    io.BytesIO(_lost_table_email()),
                    "pdf-source.eml",
                    "message/rfc822",
                )
            },
            headers={"X-Asset-Hub-Upload": "email-v1"},
            content_type="multipart/form-data",
        ).get_json()
        handle = uploaded["cases"][0]["source_eml"]["handle"]

        individual = client.post(
            "/api/mail-pdfs/individual",
            json={"mail_artifact_handle": handle},
        )
        assert individual.status_code == 201
        individual_pdf = client.get(individual.get_json()["download_url"])
        assert individual_pdf.status_code == 200
        assert individual_pdf.data.startswith(b"%PDF")
        assert individual_pdf.headers["Cache-Control"] == "private, no-store"
        individual_pdf.close()

        batch = client.post(
            "/api/mail-pdfs/batches",
            json={
                "mail_artifact_handles": [handle],
                "pages_per_mail": 2,
                "overflow_policy": "fail",
                "batch_name": "GN2140826",
            },
        )
        assert batch.status_code == 201
        batch_payload = batch.get_json()
        assert batch_payload["output_name"] == "chungtu_GN2140826.pdf"
        assert batch_payload["items"][0]["output_pages"] == 2
        merged = client.get(batch_payload["merged_download_url"])
        item = client.get(batch_payload["items"][0]["download_url"])
        assert merged.data.startswith(b"%PDF")
        assert "chungtu_GN2140826.pdf" in merged.headers["Content-Disposition"]
        assert item.data.startswith(b"%PDF")
        assert merged.headers["Cache-Control"] == "private, no-store"
        assert item.headers["Cache-Control"] == "private, no-store"
        merged.close()
        item.close()

        cleared = client.post(
            "/api/test-data/clear",
            json={"confirm": "CLEAR_TEST_DATA"},
            headers={"X-Asset-Hub-Action": "clear-test-data-v1"},
        )
        assert cleared.status_code == 200
        assert cleared.get_json()["tran_managed_output_count"] == 3
    finally:
        app.extensions["asset_hub"]["repository"].close()
