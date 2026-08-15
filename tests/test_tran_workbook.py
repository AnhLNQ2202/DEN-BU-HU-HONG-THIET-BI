"""Synthetic template tests for the TranNNB workbook exporter."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from asset_compensation.adapters import (
    TRAN_SENT_HEADERS,
    TRAN_YEAR_HEADERS,
    OutputExistsError,
    TranWorkbookAdapter,
    TranWorkbookError,
    remaining_value_excel_formula,
)
from asset_compensation.domain import (
    CompensationAsset,
    CompensationStatus,
    DepreciationGroup,
    ReferenceStatus,
)
from asset_compensation.services import (
    CompensationService,
    TranAssetRequest,
    TranResolution,
)

_BUILT_IN_TRAN_TEMPLATE = (
    Path(__file__).resolve().parents[1]
    / "src"
    / "asset_compensation"
    / "templates"
    / "tran_compensation_template.xlsx"
)


def _template(path: Path) -> Path:
    workbook = Workbook()
    year = workbook.active
    year.title = "2026"
    year.cell(1, 1, "Synthetic template marker")
    for column, header in enumerate(TRAN_YEAR_HEADERS, start=1):
        year.cell(3, column, header)
        year.cell(4, column).fill = PatternFill("solid", fgColor="FFF2CC")
        year.cell(4, column).font = Font(name="Arial", size=10)
        year.cell(4, column).alignment = Alignment(vertical="center")
    history = [
        1,
        date(2025, 12, 1),
        "OLD-001",
        "MOU00001",
        "Historical synthetic row",
        "history.user",
        date(2024, 1, 1),
        date(2025, 12, 1),
        600_000,
        60_000,
        30_000,
        90_000,
        23,
        "Tool",
        "VNG",
        "0603",
        "000",
        "01",
        25,
        0.05,
        None,
        "Synthetic historical row",
    ]
    for column, value in enumerate(history, start=1):
        year.cell(4, column, value)
    sent = workbook.create_sheet("Sent out")
    sent.append(list(TRAN_SENT_HEADERS))
    sent.append(["STALE"] * 15)
    sent.append(["STALE-TOTAL"] * 15)
    workbook.create_sheet("Unrelated")["A1"] = "Preserve me"
    workbook.save(path)
    workbook.close()
    return path


def _resolution(
    *,
    tag: str = "MOU10001",
    cost: int = 1_000_000,
    start: date = date(2024, 11, 1),
    lost: date = date(2026, 1, 1),
) -> TranResolution:
    asset = CompensationAsset(
        tag_number=tag,
        asset_name=f"Synthetic {tag}",
        domain="demo.user",
        lost_date=lost,
        cost=cost,
        start_date=start,
        asset_number="SYN-001",
        book="Tool",
        entity="VNG",
        cost_center="0603",
        product_code="000",
        location="01",
        physical=True,
        lookup_status=ReferenceStatus.MATCHED,
    )
    preview = CompensationService().preview(asset)
    return TranResolution(
        request=TranAssetRequest(
            tag_number=tag,
            asset_name=asset.asset_name,
            domain=asset.domain,
            lost_date=lost,
        ),
        asset=asset,
        preview=preview,
        notes=("Synthetic provenance note.",),
        issues=(),
        fa_status=ReferenceStatus.MATCHED,
        classification_status=ReferenceStatus.MATCHED,
    )


def test_full_formula_lists_past_current_and_future_years() -> None:
    formula = remaining_value_excel_formula(9, DepreciationGroup.FOUR_YEAR, 14)

    assert formula == (
        "=ROUND(I9*(50%*(0/12) + 30%*(10/12) + 10% + 10%),0)"
    )
    assert remaining_value_excel_formula(9, DepreciationGroup.SIX_YEAR, 72) == (
        "=ROUND(I9*10%,0)"
    )


def test_export_appends_year_rows_and_rebuilds_request_only_sent_out(tmp_path: Path) -> None:
    template = _template(tmp_path / "synthetic-template.xlsx")
    before_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    output = tmp_path / "outputs" / "tran-result.xlsx"
    calculated = _resolution()
    exempt = _resolution(
        tag="MOU10002",
        cost=499_999,
        start=date(2025, 1, 1),
        lost=date(2026, 1, 1),
    )

    result = TranWorkbookAdapter(template).export(
        [calculated, exempt], output, processing_date=date(2026, 1, 15)
    )

    assert result.request_rows == (5, 6)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before_hash
    workbook = load_workbook(output, data_only=False)
    try:
        year = workbook["2026"]
        sent = workbook["Sent out"]
        assert year["A1"].value == "Synthetic template marker"
        assert workbook["Unrelated"]["A1"].value == "Preserve me"
        assert year["D4"].value == "MOU00001"
        assert year["D5"].value == "MOU10001"
        assert year["J5"].value == (
            "=ROUND(I5*(50%*(0/12) + 30%*(10/12) + 10% + 10%),0)"
        )
        assert year["K5"].value == "=ROUND(I5*T5,0)"
        assert year["L5"].value == "=ROUND(SUM(J5:K5),0)"
        assert year["M5"].value == "=ROUND(DAYS360(G5,H5,TRUE)/30,0)"
        assert year["J6"].value == "Không tính đền bù"
        assert tuple(sent.cell(1, column).value for column in range(1, 16)) == (
            TRAN_SENT_HEADERS
        )
        assert sent.max_row == 4
        assert sent["A2"].value == "MOU10001"
        assert sent["G2"].data_type == "n"
        assert sent["G2"].value == calculated.preview.remaining_value
        assert sent["G3"].value == "Không tính đền bù"
        assert sent["G4"].value == "=SUM(G2:G3)"
        assert sent["I4"].value == "=SUM(I2:I3)"
        assert sent["A1"].font.name == "Arial"
        assert sent["A1"].fill.fgColor.theme == 4
        assert sent["A1"].border.left.style == "thin"
        assert sent["D2"].number_format == "dd/mm/yyyy"
        assert sent["F2"].number_format == "#,##0"
        assert sent["B4"].font.bold is True
    finally:
        workbook.close()


def test_export_never_overwrites_template_or_existing_output(tmp_path: Path) -> None:
    template = _template(tmp_path / "synthetic-template.xlsx")
    adapter = TranWorkbookAdapter(template)
    with pytest.raises(TranWorkbookError, match="never be overwritten"):
        adapter.export([_resolution()], template, processing_date=date(2026, 1, 1))

    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing")
    with pytest.raises(OutputExistsError):
        adapter.export([_resolution()], output, processing_date=date(2026, 1, 1))


def test_export_blocks_items_that_still_need_review(tmp_path: Path) -> None:
    template = _template(tmp_path / "synthetic-template.xlsx")
    resolution = _resolution(tag="QQQ10001")
    assert resolution.preview is not None
    assert resolution.preview.status is CompensationStatus.NEEDS_REVIEW

    with pytest.raises(TranWorkbookError, match="requires manual review"):
        TranWorkbookAdapter(template).export(
            [resolution], tmp_path / "blocked.xlsx", processing_date=date(2026, 1, 1)
        )


def test_export_lists_verified_nonphysical_item_without_inventing_reference_values(
    tmp_path: Path,
) -> None:
    template = _template(tmp_path / "synthetic-template.xlsx")
    asset = CompensationAsset(
        tag_number="SRV10001",
        asset_name="Synthetic cloud service",
        domain="demo.user",
        lost_date=date(2026, 1, 1),
        cost=None,
        start_date=None,
        physical=False,
        lookup_status=None,
    )
    resolution = TranResolution(
        request=TranAssetRequest(
            tag_number=asset.tag_number,
            asset_name=asset.asset_name,
            domain=asset.domain,
            lost_date=asset.lost_date,
        ),
        asset=asset,
        preview=CompensationService().preview(asset),
        notes=("Verified nonphysical in synthetic Define.",),
        issues=(),
        fa_status=ReferenceStatus.NOT_FOUND,
        classification_status=ReferenceStatus.MATCHED,
    )
    output = tmp_path / "nonphysical.xlsx"

    TranWorkbookAdapter(template).export(
        [resolution], output, processing_date=date(2026, 1, 1)
    )

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["2026"]["G5"].value is None
        assert workbook["2026"]["I5"].value is None
        assert workbook["2026"]["J5"].value == "Không áp dụng"
        assert workbook["Sent out"]["G2"].value == "Không áp dụng"
    finally:
        workbook.close()


def test_sanitized_built_in_template_is_export_ready_and_contains_no_case_data(
    tmp_path: Path,
) -> None:
    template = load_workbook(_BUILT_IN_TRAN_TEMPLATE, data_only=False)
    try:
        year = template["2026"]
        assert tuple(year.cell(3, column).value for column in range(1, 23)) == (
            TRAN_YEAR_HEADERS
        )
        assert all(year.cell(4, column).value is None for column in range(1, 23))
        assert tuple(
            template["Sent out"].cell(1, column).value for column in range(1, 16)
        ) == TRAN_SENT_HEADERS
    finally:
        template.close()

    output = tmp_path / "built-in-tran-output.xlsx"
    result = TranWorkbookAdapter(_BUILT_IN_TRAN_TEMPLATE).export(
        [_resolution()], output, processing_date=date(2026, 1, 15)
    )

    assert result.request_rows == (4,)
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["2026"]["D4"].value == "MOU10001"
        assert workbook["2026"]["J4"].data_type == "f"
        assert workbook["Sent out"]["A2"].value == "MOU10001"
        assert workbook["Sent out"]["B3"].value == "Total:"
    finally:
        workbook.close()
